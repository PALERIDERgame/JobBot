from __future__ import annotations

import logging
import re
import shutil
import subprocess
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from config import sanitize_filename
from database import Job
from document_tailoring import DocumentTailoringPayload
from match_scorer import MatchScore
from resume_parser import ResumeData, ResumeParagraphTemplate, ResumeWorkEntry

try:
    import pdfplumber
except ImportError:  # pragma: no cover
    pdfplumber = None


LOGGER = logging.getLogger(__name__)
_STOPWORDS = {
    "and", "the", "for", "with", "from", "that", "this", "your", "will", "role", "team", "work", "into",
    "their", "about", "across", "have", "has", "our", "you", "job", "position", "new", "york", "experience",
    "using", "within", "required", "preferred", "strong",
}
_INFERENCE_RULES = (
    {
        "skill": "Process Improvement",
        "triggers": ("process improvement", "continuous improvement", "workflow optimization"),
        "evidence": ("streamlined", "optimized", "implemented", "improved", "efficiency"),
    },
    {
        "skill": "Stakeholder Management",
        "triggers": ("stakeholder management", "stakeholder engagement"),
        "evidence": ("community leaders", "clients", "coalition partners", "donors", "executive leadership", "referral network"),
    },
    {
        "skill": "Cross-Functional Collaboration",
        "triggers": ("cross-functional", "cross functional"),
        "evidence": ("collaborated", "coordinated", "worked with", "team of", "freelance"),
    },
    {
        "skill": "Dashboard Reporting",
        "triggers": ("dashboard", "reporting", "kpi"),
        "evidence": ("dashboard", "reporting", "metrics", "analysis", "data"),
    },
    {
        "skill": "Vendor Management",
        "triggers": ("vendor management", "agency management"),
        "evidence": ("vendors", "graphic designers", "photographers", "videographers", "freelance"),
    },
    {
        "skill": "Email Marketing",
        "triggers": ("email marketing", "lifecycle marketing"),
        "evidence": ("mailchimp", "salesforce marketing cloud", "email messaging"),
    },
    {
        "skill": "CRM Management",
        "triggers": ("crm", "crm management"),
        "evidence": ("crm", "salesforce"),
    },
    {
        "skill": "SEO",
        "triggers": ("seo", "search engine optimization"),
        "evidence": ("seo", "google analytics", "wordpress"),
    },
    {
        "skill": "Conversion Optimization",
        "triggers": ("conversion", "cro", "conversion optimization"),
        "evidence": ("testing", "optimization", "performance", "landing page", "website", "campaign"),
    },
    {
        "skill": "Digital Merchandising",
        "triggers": ("merchandising", "product assortment", "product strategy"),
        "evidence": ("campaign calendar", "product launch", "audience strategy", "offerings", "promotions"),
    },
    {
        "skill": "Customer Retention",
        "triggers": ("retention", "repeat", "loyalty", "lifecycle"),
        "evidence": ("email marketing", "supporters", "crm", "salesforce marketing cloud", "mailchimp"),
    },
    {
        "skill": "Performance Marketing",
        "triggers": ("performance marketing", "customer acquisition", "paid media", "digital advertising"),
        "evidence": ("paid media", "digital ad", "campaign", "lead generation", "advertising"),
    },
    {
        "skill": "Marketing Strategy",
        "triggers": ("marketing strategy", "strategic growth", "growth initiatives"),
        "evidence": ("strategy", "campaign", "growth", "planning", "marketing"),
    },
    {
        "skill": "Lead Generation",
        "triggers": ("lead generation", "pipeline", "closed-won", "closed won"),
        "evidence": ("lead generation", "leads", "pipeline", "supporters", "donors", "campaign"),
    },
    {
        "skill": "Sales Enablement",
        "triggers": ("sales support", "sales enablement", "sales leaders", "sales organizations"),
        "evidence": ("press releases", "white papers", "messaging", "campaign", "clients"),
    },
    {
        "skill": "Product Marketing",
        "triggers": ("product marketing", "value propositions", "positioning", "market segment"),
        "evidence": ("brand", "messaging", "creative", "content", "press releases", "white papers"),
    },
    {
        "skill": "ROI Reporting",
        "triggers": ("roi", "kpi", "dashboard", "resource allocation"),
        "evidence": ("dashboard", "reporting", "metrics", "analysis", "resource allocation", "data"),
    },
    {
        "skill": "Audience Segmentation",
        "triggers": ("audience segmentation", "segmentation", "audience targeting"),
        "evidence": ("demographics", "voter outreach", "lead generation", "supporters"),
    },
    {
        "skill": "Program Management",
        "triggers": ("program management",),
        "evidence": ("programming", "planned and ran", "oversaw", "managed a staff"),
    },
)
_INFERENCE_BLOCKLIST = (
    "certification", "certificate", "licensed", "license", "licensure", "pmp", "cpa", "rn",
)


@dataclass(slots=True)
class GeneratedDocs:
    output_dir: Path
    resume_pdf_path: Path
    cover_letter_path: Path
    resume_docx_path: Path = Path()
    status: str = "generated"
    error_message: str = ""
    pdf_exporter_used: str = ""
    page_fit_attempts: int = 0


@dataclass(slots=True)
class SkillCandidate:
    name: str
    confidence: float
    direct: bool


@dataclass(slots=True)
class ResumeFitOptions:
    compact_level: int = 0
    per_role_trim_rounds: int = 0
    global_trim_count: int = 0
    key_skills_limit: int | None = None
    phase: str = "initial"


class DocumentGenerator:
    def generate(
        self,
        base_output_dir: Path,
        job: Job,
        resume: ResumeData,
        score: MatchScore,
        *,
        ai_notes: str = "",
        tailoring_payload: DocumentTailoringPayload | None = None,
        progress_callback=None,
    ) -> GeneratedDocs:
        def report(stage: str, message: str, progress: int) -> None:
            if progress_callback:
                progress_callback(stage, message, progress)

        report("starting", "Preparing output folder...", 1)
        folder_name = f"{sanitize_filename(job.employer)}_{datetime.now().strftime('%Y-%m-%d')}_{job.id[:8]}"
        output_dir = base_output_dir / folder_name
        output_dir.mkdir(parents=True, exist_ok=True)
        staging_dir = output_dir / "_staging"
        if staging_dir.exists():
            shutil.rmtree(staging_dir, ignore_errors=True)
        staging_dir.mkdir(parents=True, exist_ok=True)

        staged_resume_pdf_path = staging_dir / "resume.pdf"
        staged_resume_docx_path = staging_dir / "resume.docx"
        staged_cover_letter_path = staging_dir / "cover_letter.txt"
        resume_pdf_path = output_dir / "resume.pdf"
        resume_docx_path = output_dir / "resume.docx"
        cover_letter_path = output_dir / "cover_letter.txt"
        source_path = Path(resume.source_path) if resume.source_path else None

        status = "generated"
        error_message = ""
        pdf_exporter_used = ""
        page_fit_attempts = 0
        if source_path and source_path.suffix.lower() == ".docx" and source_path.exists():
            fit_attempts = self._build_fit_attempts(resume, job)
            fit_succeeded = False
            try:
                for attempt_idx, attempt in enumerate(fit_attempts, 1):
                    page_fit_attempts = attempt_idx
                    report("tailoring_resume", self._fit_attempt_message(attempt), 3)
                    self._build_docx_resume(
                        staged_resume_docx_path,
                        job,
                        resume,
                        source_path,
                        fit_options=attempt,
                        tailoring_payload=tailoring_payload,
                    )
                    try:
                        report("exporting_pdf", "Exporting resume PDF...", 4)
                        pdf_exporter_used = self._export_docx_to_pdf(staged_resume_docx_path, staged_resume_pdf_path)
                        report("validating_pages", "Validating final page count...", 5)
                        if self._count_pdf_pages(staged_resume_pdf_path) <= 2:
                            fit_succeeded = True
                            break
                    except Exception as exc:
                        status = "generated_docx_only"
                        error_message = str(exc)
                        staged_resume_pdf_path = Path()
                        fit_succeeded = True
                        break
                if not fit_succeeded:
                    raise ValueError("JobBot could not compress the tailored resume to 2 pages.")
            except Exception as exc:
                if "compress the tailored resume to 2 pages" in str(exc).lower():
                    shutil.rmtree(staging_dir, ignore_errors=True)
                    return GeneratedDocs(
                        output_dir=output_dir,
                        resume_pdf_path=Path(),
                        cover_letter_path=Path(),
                        resume_docx_path=Path(),
                        status="failed",
                        error_message=str(exc),
                        pdf_exporter_used=pdf_exporter_used,
                        page_fit_attempts=page_fit_attempts,
                    )
                raise
        else:
            report("tailoring_resume", "Generating resume PDF...", 3)
            self._build_resume_pdf(staged_resume_pdf_path, job, resume, score, ai_notes=ai_notes)
            resume_docx_path = Path()
            staged_resume_docx_path = Path()

        report("writing_cover_letter", "Writing cover letter...", 6)
        self._build_cover_letter(
            staged_cover_letter_path,
            job,
            resume,
            score,
            ai_notes=ai_notes,
            tailoring_payload=tailoring_payload,
        )
        resume_docx_path = self._promote_staged_file(staged_resume_docx_path, resume_docx_path)
        resume_pdf_path = self._promote_staged_file(staged_resume_pdf_path, resume_pdf_path)
        cover_letter_path = self._promote_staged_file(staged_cover_letter_path, cover_letter_path)
        shutil.rmtree(staging_dir, ignore_errors=True)
        LOGGER.info("Generated documents for %s at %s", job.id, output_dir)
        return GeneratedDocs(
            output_dir=output_dir,
            resume_pdf_path=resume_pdf_path,
            cover_letter_path=cover_letter_path,
            resume_docx_path=resume_docx_path,
            status=status,
            error_message=error_message,
            pdf_exporter_used=pdf_exporter_used,
            page_fit_attempts=page_fit_attempts,
        )

    def _build_docx_resume(
        self,
        path: Path,
        job: Job,
        resume: ResumeData,
        source_docx_path: Path,
        *,
        fit_options: ResumeFitOptions | None = None,
        tailoring_payload: DocumentTailoringPayload | None = None,
    ) -> None:
        try:
            from docx import Document as DocxDocument
            from docx.text.paragraph import Paragraph
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-docx is required to generate DOCX resumes") from exc

        document = DocxDocument(str(source_docx_path))
        paragraphs = list(document.paragraphs)
        source_paragraphs = list(document.paragraphs)
        work_heading_idx = self._find_heading_index(paragraphs, "WORK EXPERIENCE")
        education_heading_idx = self._find_heading_index(paragraphs, "EDUCATION")
        key_skills_heading_idx = self._find_heading_index(paragraphs, "KEY SKILLS")
        if work_heading_idx < 0 or education_heading_idx < 0 or key_skills_heading_idx < 0:
            raise ValueError("DOCX resume is missing WORK EXPERIENCE, EDUCATION, or KEY SKILLS headings")

        education_heading = paragraphs[education_heading_idx]
        key_skills_heading = paragraphs[key_skills_heading_idx]
        work_body = paragraphs[work_heading_idx + 1:education_heading_idx]
        key_skills_body = paragraphs[key_skills_heading_idx + 1:]

        for paragraph in work_body:
            self._remove_paragraph(paragraph)
        for paragraph in key_skills_body:
            self._remove_paragraph(paragraph)

        base_entries = tailoring_payload.work_entries if tailoring_payload and tailoring_payload.work_entries else None
        for entry in self._tailor_work_entries(resume, job, fit_options=fit_options, source_entries=base_entries):
            role_template = self._paragraph_from_template(source_paragraphs, entry.get("role_template"))
            date_template = self._paragraph_from_template(source_paragraphs, entry.get("date_template"))
            bullet_templates = [
                template_paragraph
                for template_paragraph in (
                    self._paragraph_from_template(source_paragraphs, template)
                    for template in entry.get("bullet_templates", [])
                )
                if template_paragraph is not None
            ]
            if role_template is None:
                raise ValueError("DOCX resume is missing a reusable role paragraph template")
            self._clone_paragraph_before(education_heading, role_template, str(entry["role_line"]))
            if entry["date_line"]:
                self._clone_paragraph_before(education_heading, date_template or role_template, str(entry["date_line"]))
            if not bullet_templates:
                raise ValueError("DOCX resume is missing reusable bullet paragraph templates")
            for idx, bullet in enumerate(entry["bullets"]):
                template = bullet_templates[min(idx, len(bullet_templates) - 1)]
                self._clone_paragraph_before(education_heading, template, str(bullet))

        key_skill_templates = [
            template_paragraph
            for template_paragraph in (
                self._paragraph_from_template(source_paragraphs, template)
                for template in resume.key_skills_templates
            )
            if template_paragraph is not None
        ]
        if not key_skill_templates:
            key_skill_templates = [self._first_nonempty_paragraph(key_skills_body) or key_skills_heading]

        tailored_skills = self._tailor_key_skills(
            resume,
            job,
            base_skills=tailoring_payload.key_skills if tailoring_payload and tailoring_payload.key_skills else None,
            limit=fit_options.key_skills_limit if fit_options else None,
        )
        skill_lines = self._render_key_skills_lines(tailored_skills, line_count=max(1, len(key_skill_templates)))
        anchor = key_skills_heading
        for idx, skill_line in enumerate(skill_lines):
            template = key_skill_templates[min(idx, len(key_skill_templates) - 1)]
            anchor = self._clone_paragraph_after(anchor, template, skill_line)

        document.save(str(path))

    def _build_resume_pdf(self, path: Path, job: Job, resume: ResumeData, score: MatchScore, *, ai_notes: str) -> None:
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError:
            self._write_fallback_pdf(path, job, resume, score)
            return

        styles = getSampleStyleSheet()
        contact_line = " | ".join(part for part in [resume.email, resume.phone] if part)
        clean_employer = self._display_employer(job.employer)
        clean_summary = self._clean_summary(resume.summary)
        skills_text = ", ".join(self._tailor_key_skills(resume, job)[:12]) or "Relevant skills were not confidently extracted from the resume source."
        experience_lines = resume.experience_lines[:6] or ["Relevant experience highlights were not confidently extracted from the resume source."]
        doc = SimpleDocTemplate(str(path), pagesize=letter)
        story = [
            Paragraph(resume.name, styles["Title"]),
            Paragraph(contact_line or "Contact details unavailable", styles["Normal"]),
            Spacer(1, 12),
            Paragraph("Professional Summary", styles["Heading2"]),
            Paragraph(clean_summary, styles["BodyText"]),
            Spacer(1, 12),
            Paragraph("Target Role", styles["Heading2"]),
            Paragraph(f"{job.title} at {clean_employer}", styles["BodyText"]),
            Spacer(1, 12),
            Paragraph("Relevant Skills", styles["Heading2"]),
            Paragraph(skills_text, styles["BodyText"]),
            Spacer(1, 12),
            Paragraph("Relevant Experience Highlights", styles["Heading2"]),
        ]
        for line in experience_lines:
            story.append(Paragraph(f"- {line}", styles["BodyText"]))
        story.extend(
            [
                Spacer(1, 12),
                Paragraph("Match Notes", styles["Heading2"]),
                Paragraph(f"Match score: {score.score if score.score is not None else 'N/A'}", styles["BodyText"]),
                Paragraph(self._clean_rationale(score.rationale), styles["BodyText"]),
            ]
        )
        if ai_notes:
            story.extend(
                [
                    Spacer(1, 12),
                    Paragraph("AI Tailoring Notes", styles["Heading2"]),
                    Paragraph(ai_notes.replace("\n", "<br/>"), styles["BodyText"]),
                ]
            )
        doc.build(story)

    def _build_cover_letter(
        self,
        path: Path,
        job: Job,
        resume: ResumeData,
        score: MatchScore,
        *,
        ai_notes: str,
        tailoring_payload: DocumentTailoringPayload | None = None,
    ) -> None:
        clean_employer = self._display_employer(job.employer)
        signoff_name = self._signoff_name(resume)
        if tailoring_payload and tailoring_payload.cover_letter_text:
            letter_text = self._normalize_output_text(tailoring_payload.cover_letter_text).strip()
            cover_letter_ok, _ = self._cover_letter_quality_issue(letter_text)
            if not cover_letter_ok:
                letter_text = self._compose_cover_letter(job, resume, score, clean_employer, signoff_name)
            elif "sincerely" not in letter_text.lower():
                letter_text = "\n".join([letter_text, "", f"Sincerely,\n{signoff_name}"])
        elif ai_notes:
            letter_text = "\n".join(
                [
                    f"Dear Hiring Team at {clean_employer},",
                    "",
                    self._normalize_output_text(ai_notes),
                    "",
                    f"Sincerely,\n{signoff_name}",
                ]
            )
        else:
            letter_text = self._compose_cover_letter(job, resume, score, clean_employer, signoff_name)
        path.write_text(letter_text, encoding="utf-8")

    def _export_docx_to_pdf(self, docx_path: Path, pdf_path: Path) -> str:
        """Export docx to PDF. Returns the exporter name used ('word_com' or 'libreoffice')."""
        word_error = None
        try:
            import win32com.client  # type: ignore[import-not-found]
        except ImportError:
            win32com = None
        else:
            win32com = win32com.client

        if win32com is not None:
            word = None
            document = None
            try:
                LOGGER.info("Attempting Word PDF export for %s", docx_path)
                word = win32com.DispatchEx("Word.Application")
                word.Visible = False
                document = word.Documents.Open(str(docx_path.resolve()))
                document.SaveAs(str(pdf_path.resolve()), FileFormat=17)
                LOGGER.info("Word PDF export succeeded for %s", docx_path)
                return "word_com"
            except Exception as exc:
                word_error = str(exc)
                LOGGER.warning("Word PDF export failed for %s: %s", docx_path, exc)
            finally:
                if document is not None:
                    document.Close(False)
                if word is not None:
                    word.Quit()

        libreoffice_errors: list[str] = []
        for soffice in self._libreoffice_candidates():
            LOGGER.info("Attempting LibreOffice PDF export for %s via %s", docx_path, soffice)
            completed = subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)],
                capture_output=True,
                text=True,
                check=False,
            )
            if completed.returncode == 0 and pdf_path.exists():
                LOGGER.info("LibreOffice PDF export succeeded for %s via %s", docx_path, soffice)
                return "libreoffice"
            error_text = (completed.stderr or completed.stdout or "LibreOffice PDF export failed").strip()
            libreoffice_errors.append(f"{soffice}: {error_text}")
            LOGGER.warning("LibreOffice PDF export failed for %s via %s: %s", docx_path, soffice, error_text)

        if libreoffice_errors:
            prefix = f"Word automation failed: {word_error}. " if word_error else ""
            raise RuntimeError(prefix + "LibreOffice conversion failed: " + " | ".join(libreoffice_errors))
        if word_error:
            raise RuntimeError(f"Word automation failed: {word_error}. No PDF exporter found.")
        raise RuntimeError("No PDF exporter found for resume.pdf export.")

    @staticmethod
    def _libreoffice_candidates() -> list[str]:
        candidates: list[str] = []
        path_candidate = shutil.which("soffice")
        if path_candidate:
            candidates.append(path_candidate)
        for candidate in (
            Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
            Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
        ):
            if candidate.exists():
                candidates.append(str(candidate))
        ordered: list[str] = []
        seen: set[str] = set()
        for candidate in candidates:
            normalized = candidate.lower()
            if normalized in seen:
                continue
            seen.add(normalized)
            ordered.append(candidate)
        return ordered

    @staticmethod
    def _count_pdf_pages(pdf_path: Path) -> int:
        if pdfplumber is None:  # pragma: no cover
            raise RuntimeError("pdfplumber is required to validate PDF page count.")
        with pdfplumber.open(pdf_path) as pdf:
            return len(pdf.pages)

    @staticmethod
    def _find_heading_index(paragraphs: list[object], heading_text: str) -> int:
        target = heading_text.strip().lower()
        for idx, paragraph in enumerate(paragraphs):
            if getattr(paragraph, "text", "").strip().lower() == target:
                return idx
        return -1

    @staticmethod
    def _is_date_text(text: str) -> bool:
        return bool(re.match(r"^(JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST|SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)", text.strip(), re.IGNORECASE))

    def _find_work_template(self, paragraphs: list[object], *, kind: str):
        for paragraph in paragraphs:
            text = getattr(paragraph, "text", "").strip()
            if not text:
                continue
            style_name = (getattr(getattr(paragraph, "style", None), "name", "") or "").lower()
            if kind == "role" and not self._is_date_text(text) and "bullet" not in style_name:
                return paragraph
            if kind == "date" and self._is_date_text(text):
                return paragraph
            if kind == "bullet" and ("bullet" in style_name or text.startswith(("•", "-", "*"))):
                return paragraph
        return paragraphs[0] if paragraphs else None

    @staticmethod
    def _first_nonempty_paragraph(paragraphs: list[object]):
        for paragraph in paragraphs:
            if getattr(paragraph, "text", "").strip():
                return paragraph
        return None

    @staticmethod
    def _paragraph_from_template(paragraphs, template: ResumeParagraphTemplate | None):
        if template is None:
            return None
        try:
            return paragraphs[template.paragraph_index]
        except IndexError:
            return None

    def _clone_paragraph_before(self, anchor, template, text: str):
        from docx.text.paragraph import Paragraph

        cloned = deepcopy(template._p)
        anchor._p.addprevious(cloned)
        paragraph = Paragraph(cloned, anchor._parent)
        self._apply_paragraph_template(paragraph, template, text)
        if self._is_bullet_template(template):
            self._apply_generated_bullet_numbering(paragraph)
        return paragraph

    def _clone_paragraph_after(self, anchor, template, text: str):
        from docx.text.paragraph import Paragraph

        cloned = deepcopy(template._p)
        anchor._p.addnext(cloned)
        paragraph = Paragraph(cloned, anchor._parent)
        self._apply_paragraph_template(paragraph, template, text)
        return paragraph

    def _apply_paragraph_template(self, paragraph, template, text: str) -> None:
        from docx.shared import Pt

        for child in list(paragraph._p):
            if child.tag.endswith("}pPr"):
                continue
            paragraph._p.remove(child)
        run = paragraph.add_run(text)
        template_run = template.runs[0] if template is not None and getattr(template, "runs", None) else None
        if template_run is not None:
            font = run.font
            source_font = template_run.font
            font.bold = source_font.bold
            font.italic = source_font.italic
            font.name = source_font.name
            font.size = source_font.size
            font.underline = source_font.underline
            if self._is_bullet_template(template):
                if font.size is not None:
                    font.size = max(font.size - 6350, 91440)
                else:
                    font.size = Pt(10)

    @staticmethod
    def _is_bullet_template(template) -> bool:
        style_name = (getattr(getattr(template, "style", None), "name", "") or "").lower()
        p_pr = getattr(template._p, "pPr", None)
        has_numbering = bool(p_pr is not None and getattr(p_pr, "numPr", None) is not None)
        paragraph_format = getattr(template, "paragraph_format", None)
        first_line_indent = getattr(paragraph_format, "first_line_indent", None)
        left_indent = getattr(paragraph_format, "left_indent", None)
        has_hanging_indent = bool(
            first_line_indent is not None
            and left_indent is not None
            and first_line_indent < 0
            and left_indent > 0
        )
        return has_numbering or "bullet" in style_name or has_hanging_indent

    def _apply_generated_bullet_numbering(self, paragraph) -> None:
        try:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError:  # pragma: no cover
            return

        numbering_info = self._paragraph_numbering_info(paragraph)
        if numbering_info is None:
            return
        source_num_id, ilvl = numbering_info

        numbering_part = getattr(paragraph.part, "numbering_part", None)
        numbering_root = getattr(numbering_part, "element", None)
        if numbering_root is None:
            return

        cache = getattr(numbering_part, "_jobbot_bullet_num_cache", None)
        if cache is None:
            cache = {}
            setattr(numbering_part, "_jobbot_bullet_num_cache", cache)
        cache_key = (source_num_id, ilvl, 12)
        if cache_key not in cache:
            num_xpath = f'.//*[local-name()="num" and @*[local-name()="numId"]="{source_num_id}"]'
            num_nodes = numbering_root.xpath(num_xpath)
            if not num_nodes:
                return
            source_num = num_nodes[0]
            abstract_num_id_element = source_num.find(qn("w:abstractNumId"))
            if abstract_num_id_element is None:
                return
            source_abstract_num_id = abstract_num_id_element.get(qn("w:val"))
            if not source_abstract_num_id:
                return
            abstract_xpath = f'.//*[local-name()="abstractNum" and @*[local-name()="abstractNumId"]="{source_abstract_num_id}"]'
            abstract_nodes = numbering_root.xpath(abstract_xpath)
            if not abstract_nodes:
                return
            source_abstract = abstract_nodes[0]

            existing_abstract_ids = [
                int(node.get(qn("w:abstractNumId")))
                for node in numbering_root.xpath('.//*[local-name()="abstractNum"]')
                if node.get(qn("w:abstractNumId")) is not None
            ]
            existing_num_ids = [
                int(node.get(qn("w:numId")))
                for node in numbering_root.xpath('.//*[local-name()="num"]')
                if node.get(qn("w:numId")) is not None
            ]
            new_abstract_id = (max(existing_abstract_ids) + 1) if existing_abstract_ids else 100
            new_num_id = (max(existing_num_ids) + 1) if existing_num_ids else 100

            cloned_abstract = deepcopy(source_abstract)
            cloned_abstract.set(qn("w:abstractNumId"), str(new_abstract_id))
            for lvl in cloned_abstract.xpath('.//*[local-name()="lvl"]'):
                r_pr = lvl.find(qn("w:rPr"))
                if r_pr is None:
                    r_pr = OxmlElement("w:rPr")
                    lvl.append(r_pr)
                for tag_name in ("w:sz", "w:szCs"):
                    existing = r_pr.find(qn(tag_name))
                    if existing is None:
                        existing = OxmlElement(tag_name)
                        r_pr.append(existing)
                    existing.set(qn("w:val"), "12")

            cloned_num = deepcopy(source_num)
            cloned_num.set(qn("w:numId"), str(new_num_id))
            cloned_abstract_ref = cloned_num.find(qn("w:abstractNumId"))
            if cloned_abstract_ref is not None:
                cloned_abstract_ref.set(qn("w:val"), str(new_abstract_id))

            numbering_root.append(cloned_abstract)
            numbering_root.append(cloned_num)
            cache[cache_key] = new_num_id

        p_pr = getattr(paragraph._p, "pPr", None)
        if p_pr is None:
            p_pr = OxmlElement("w:pPr")
            paragraph._p.insert(0, p_pr)
        num_pr = getattr(p_pr, "numPr", None)
        if num_pr is None:
            num_pr = OxmlElement("w:numPr")
            p_pr.append(num_pr)
        ilvl_node = getattr(num_pr, "ilvl", None)
        if ilvl_node is None:
            ilvl_node = OxmlElement("w:ilvl")
            num_pr.append(ilvl_node)
        ilvl_node.set(qn("w:val"), str(ilvl))
        num_id_element = getattr(num_pr, "numId", None)
        if num_id_element is None:
            num_id_element = OxmlElement("w:numId")
            num_pr.append(num_id_element)
        num_id_element.set(qn("w:val"), str(cache[cache_key]))

    @staticmethod
    def _paragraph_numbering_info(paragraph) -> tuple[int, int] | None:
        def _read_num_pr(num_pr) -> tuple[int, int] | None:
            if num_pr is None:
                return None
            num_id_element = getattr(num_pr, "numId", None)
            ilvl_element = getattr(num_pr, "ilvl", None)
            if num_id_element is None:
                return None
            try:
                num_id = int(num_id_element.val)
            except (TypeError, ValueError):
                return None
            ilvl = 0
            if ilvl_element is not None:
                try:
                    ilvl = int(ilvl_element.val)
                except (TypeError, ValueError):
                    ilvl = 0
            return num_id, ilvl

        p_pr = getattr(paragraph._p, "pPr", None)
        direct = _read_num_pr(getattr(p_pr, "numPr", None) if p_pr is not None else None)
        if direct is not None:
            return direct

        style_element = getattr(getattr(paragraph, "style", None), "_element", None)
        style_p_pr = getattr(style_element, "pPr", None) if style_element is not None else None
        return _read_num_pr(getattr(style_p_pr, "numPr", None) if style_p_pr is not None else None)

    @staticmethod
    def _remove_paragraph(paragraph) -> None:
        element = paragraph._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

    @staticmethod
    def _render_key_skills_lines(skills: list[str], *, line_count: int) -> list[str]:
        if not skills:
            return ["Skills could not be extracted from the source resume."]
        if line_count <= 1:
            return [", ".join(skills)]
        chunk_size = max(1, (len(skills) + line_count - 1) // line_count)
        return [", ".join(skills[idx:idx + chunk_size]) for idx in range(0, len(skills), chunk_size)]

    def build_local_tailoring_payload(self, job: Job, resume: ResumeData, score: MatchScore, *, ai_notes: str = "") -> DocumentTailoringPayload:
        tailored_entries = self._build_resume_work_entries(job, resume)
        key_skills = self._tailor_key_skills(resume, job)
        cover_letter = self._compose_cover_letter(
            job,
            resume,
            score,
            self._display_employer(job.employer),
            self._signoff_name(resume),
        )
        return DocumentTailoringPayload(
            work_entries=tailored_entries,
            key_skills=key_skills,
            cover_letter_text=cover_letter,
            route="local",
        )

    def should_escalate_tailoring(self, job: Job, resume: ResumeData, payload: DocumentTailoringPayload) -> bool:
        issues = self._tailoring_quality_issues(payload)
        if issues:
            return True
        themes = self._job_theme_labels(job)
        title = job.title.lower()
        complexity_markers = {"head", "director", "chief", "vice president", "vp", "lead"}
        if len(themes) >= 5:
            return True
        if any(marker in title for marker in complexity_markers) and len(themes) >= 3:
            return True
        if len(job.description_full) > 3500:
            return True
        return False

    def validate_tailoring_payload(self, payload: DocumentTailoringPayload) -> tuple[bool, list[str]]:
        issues = self._tailoring_quality_issues(payload)
        return not issues, issues

    def repair_ai_tailoring_payload(
        self,
        ai_payload: DocumentTailoringPayload,
        fallback_payload: DocumentTailoringPayload,
    ) -> tuple[DocumentTailoringPayload, list[str], set[str]]:
        repaired = deepcopy(ai_payload)
        issues: list[str] = []
        failed_sections: set[str] = set()
        repaired_bullets = 0

        if len(ai_payload.work_entries) != len(fallback_payload.work_entries):
            repaired.work_entries = fallback_payload.work_entries
            issues.append("resume_structure_mismatch")
            failed_sections.add("resume")
        else:
            repaired_entries: list[ResumeWorkEntry] = []
            ai_bullets_kept = 0
            for ai_entry, fallback_entry in zip(ai_payload.work_entries, fallback_payload.work_entries):
                repaired_entry_bullets: list[str] = []
                for idx, ai_bullet in enumerate(ai_entry.bullets):
                    normalized = self._normalize_output_text(ai_bullet)
                    issue = self._bullet_quality_issue(normalized)
                    if issue:
                        issues.append(issue)
                        failed_sections.add("resume")
                        repaired_bullets += 1
                        replacement = fallback_entry.bullets[min(idx, len(fallback_entry.bullets) - 1)] if fallback_entry.bullets else normalized
                        repaired_entry_bullets.append(replacement)
                    else:
                        repaired_entry_bullets.append(normalized)
                        if idx < len(fallback_entry.bullets) and normalized != fallback_entry.bullets[idx]:
                            ai_bullets_kept += 1
                if not repaired_entry_bullets:
                    repaired_entry_bullets = list(fallback_entry.bullets)
                    issues.append(f"resume_entry_fallback:{fallback_entry.role_line}")
                    failed_sections.add("resume")
                repaired_entries.append(
                    ResumeWorkEntry(
                        role_line=fallback_entry.role_line,
                        date_line=fallback_entry.date_line,
                        bullets=repaired_entry_bullets,
                        role_template=fallback_entry.role_template,
                        date_template=fallback_entry.date_template,
                        bullet_templates=fallback_entry.bullet_templates,
                    )
                )
            repaired.work_entries = repaired_entries
            if ai_bullets_kept == 0:
                repaired.resume_ai_status = "local"
            elif repaired_bullets > 0:
                repaired.resume_ai_status = "partial"
            else:
                repaired.resume_ai_status = "accepted"

        cover_ok, cover_issue = self._cover_letter_quality_issue(ai_payload.cover_letter_text)
        if cover_ok:
            repaired.cover_letter_text = self._normalize_output_text(ai_payload.cover_letter_text).strip()
            repaired.cover_letter_ai_status = "accepted"
            repaired.cover_letter_fallback = ""
        else:
            repaired.cover_letter_text = fallback_payload.cover_letter_text
            repaired.cover_letter_ai_status = "local"
            repaired.cover_letter_fallback = "local"
            failed_sections.add("cover_letter")
            if cover_issue:
                issues.append(cover_issue)

        repaired.key_skills = ai_payload.key_skills or fallback_payload.key_skills
        repaired.rejected_bullets_repaired = repaired_bullets
        repaired.ai_repair_applied = repaired_bullets > 0 or repaired.cover_letter_ai_status != "accepted"
        repaired.route = "openai" if (
            repaired.resume_ai_status in {"accepted", "partial"} or repaired.cover_letter_ai_status == "accepted"
        ) else "fallback"
        return repaired, issues, failed_sections

    def _build_resume_work_entries(self, job: Job, resume: ResumeData, *, fit_options: ResumeFitOptions | None = None, source_entries=None):
        entries = self._tailor_work_entries(resume, job, fit_options=fit_options, source_entries=source_entries)
        built = []
        for entry in entries:
            built.append(
                ResumeWorkEntry(
                    role_line=str(entry["role_line"]),
                    date_line=str(entry["date_line"]),
                    bullets=[str(item) for item in entry["bullets"]],
                    role_template=entry.get("role_template"),
                    date_template=entry.get("date_template"),
                    bullet_templates=list(entry.get("bullet_templates", [])),
                )
            )
        return built

    def _tailoring_quality_issues(self, payload: DocumentTailoringPayload) -> list[str]:
        issues: list[str] = []
        cover_ok, cover_issue = self._cover_letter_quality_issue(payload.cover_letter_text)
        if not cover_ok and cover_issue:
            issues.append(cover_issue)
        bullets = [bullet for entry in payload.work_entries for bullet in entry.bullets]
        repeated_endings: dict[str, int] = {}
        for bullet in bullets:
            normalized = self._normalize_output_text(bullet)
            issue = self._bullet_quality_issue(normalized)
            if issue:
                issues.append(issue)
            lowered = normalized.lower()
            ending = " ".join(re.findall(r"[a-z0-9+#&'-]+", lowered)[-5:])
            if ending:
                repeated_endings[ending] = repeated_endings.get(ending, 0) + 1
        if any(count > 2 for count in repeated_endings.values()):
            issues.append("repeated_closing_phrases")
        return issues

    @staticmethod
    def _cover_letter_quality_issue(text: str) -> tuple[bool, str]:
        raw_text = str(text or "")
        normalized = DocumentGenerator._normalize_output_text(raw_text).strip()
        lowered = normalized.lower()
        if not normalized:
            return False, "cover_letter_missing"
        if lowered in {"{'type': 'string'}", '{"type":"string"}', '{"type": "string"}'}:
            return False, "cover_letter_placeholder"
        if normalized.startswith("{") and "type" in lowered and "string" in lowered and len(normalized) < 120:
            return False, "cover_letter_placeholder"
        lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
        if len(lines) <= 2 and any(line.lower().startswith("sincerely") for line in lines):
            return False, "cover_letter_signoff_only"
        if len(re.findall(r"[A-Za-z]{2,}", normalized)) < 20:
            return False, "cover_letter_too_short"
        return True, ""

    @staticmethod
    def _bullet_quality_issue(text: str) -> str:
        normalized = DocumentGenerator._normalize_output_text(text)
        lowered = normalized.lower()
        if re.search(r"\b(\w+)\s+\1\b", lowered):
            return f"duplicate_word:{normalized}"
        if re.search(r"\b(managed optimized|led trained|created delivered|planned ran|presented presented|captured cleaned)\b", lowered):
            return f"broken_conjunction:{normalized}"
        if re.search(r"\bidentify likely supporters execute\b", lowered):
            return f"missing_connector:{normalized}"
        if re.search(r"\b(both in-person|both in person)\.?$", lowered):
            return f"fragment:{normalized}"
        if len(re.findall(r"\w+", normalized)) < 5:
            return f"too_short:{normalized}"
        return ""

    def _tailor_work_entries(
        self,
        resume: ResumeData,
        job: Job,
        *,
        fit_options: ResumeFitOptions | None = None,
        source_entries=None,
    ) -> list[dict[str, object]]:
        job_terms = self._job_terms(job)
        fit_options = fit_options or ResumeFitOptions()
        changed_count = 0
        entries: list[dict[str, object]] = []
        removable_entries: list[dict[str, object]] = []
        opening_counts: dict[str, int] = {}
        work_entries = source_entries or resume.work_experience_entries
        for entry in work_entries:
            scored_bullets: list[tuple[int, int, str, str]] = []
            for idx, bullet in enumerate(entry.bullets):
                rewritten = self._tailor_bullet_text(
                    bullet.strip(),
                    job_terms,
                    job,
                    compact_level=fit_options.compact_level,
                )
                if rewritten != bullet.strip():
                    changed_count += 1
                scored_bullets.append(
                    (
                        self._bullet_relevance_score(rewritten, job_terms, job),
                        idx,
                        rewritten,
                        bullet.strip(),
                    )
                )
            if any(score > 0 for score, _idx, _bullet, _original in scored_bullets):
                scored_bullets.sort(key=lambda item: (item[0], -item[1]), reverse=True)
            guarded_bullets: list[tuple[int, int, str, str]] = []
            for score, idx, rewritten, original in scored_bullets:
                stem = self._opening_stem(rewritten)
                if stem and opening_counts.get(stem, 0) >= 2:
                    rewritten = self._tailor_bullet_text_with_mode(
                        original,
                        job_terms,
                        job,
                        compact_level=fit_options.compact_level,
                        allow_append=False,
                    )
                    stem = self._opening_stem(rewritten)
                if stem:
                    opening_counts[stem] = opening_counts.get(stem, 0) + 1
                guarded_bullets.append((score, idx, rewritten, original))
            reordered = [bullet for _score, _idx, bullet, _original in guarded_bullets]
            removable = max(0, len(reordered) - 1)
            per_role_trim = min(fit_options.per_role_trim_rounds, removable)
            if per_role_trim:
                reordered = reordered[: len(reordered) - per_role_trim]
                guarded_bullets = guarded_bullets[: len(guarded_bullets) - per_role_trim]
            entry_payload = {
                "role_line": entry.role_line,
                "date_line": entry.date_line,
                "bullets": reordered,
                "role_template": entry.role_template,
                "date_template": entry.date_template,
                "bullet_templates": entry.bullet_templates,
                "_scored_bullets": guarded_bullets,
            }
            entries.append(entry_payload)
            if len(reordered) > 1:
                removable_entries.append(entry_payload)

        remaining_global_trims = fit_options.global_trim_count
        while remaining_global_trims > 0:
            candidate_entries = [
                payload for payload in removable_entries
                if len(payload["bullets"]) > 1 and payload["_scored_bullets"]
            ]
            if not candidate_entries:
                break
            weakest_entry = min(
                candidate_entries,
                key=lambda payload: (
                    payload["_scored_bullets"][-1][0],
                    len(payload["bullets"]),
                ),
            )
            weakest_entry["bullets"].pop()
            weakest_entry["_scored_bullets"].pop()
            remaining_global_trims -= 1

        for entry in entries:
            entry.pop("_scored_bullets", None)
        if changed_count == 0 and entries:
            for entry in entries:
                if not entry["bullets"]:
                    continue
                entry["bullets"][0] = self._force_tailored_emphasis(str(entry["bullets"][0]), job_terms, job)
                break
        return entries

    def _tailor_key_skills(self, resume: ResumeData, job: Job, *, base_skills: list[str] | None = None, limit: int | None = None) -> list[str]:
        direct_skills = list(base_skills or (self._skills_from_lines(resume.key_skills_lines) or resume.skills))
        if not direct_skills:
            direct_skills = ["Skills could not be extracted from the source resume."]
        combined: list[SkillCandidate] = [SkillCandidate(skill, 1.0, True) for skill in direct_skills]
        combined.extend(self._infer_skills(resume, job))
        seen: set[str] = set()
        deduped: list[SkillCandidate] = []
        for candidate in combined:
            normalized = candidate.name.strip().lower()
            if not normalized or normalized in seen:
                continue
            seen.add(normalized)
            deduped.append(candidate)
        deduped.sort(
            key=lambda candidate: (
                self._skill_relevance_score(candidate.name, self._job_terms(job)),
                candidate.confidence,
                1 if candidate.direct else 0,
            ),
            reverse=True,
        )
        names = [candidate.name for candidate in deduped]
        if limit is not None:
            return names[: max(1, limit)]
        return names

    def _compose_cover_letter(self, job: Job, resume: ResumeData, score: MatchScore, employer: str, signoff_name: str) -> str:
        location_text = f" in {job.location}" if job.location else ""
        opening = (
            f"I am excited to apply for the {job.title} role at {employer}{location_text}. "
            f"My background combines {self._resume_value_summary(resume, job)}."
        )
        evidence = self._cover_letter_evidence_paragraph(resume, job)
        closing = self._cover_letter_closing(job, resume, score, employer)
        return "\n".join(
            [
                f"Dear Hiring Team at {employer},",
                "",
                opening,
                "",
                evidence,
                "",
                closing,
                "",
                f"Sincerely,\n{signoff_name}",
            ]
        )

    def _cover_letter_evidence_paragraph(self, resume: ResumeData, job: Job) -> str:
        evidence_lines = self._select_cover_letter_evidence(resume, job)
        if not evidence_lines:
            return (
                f"In prior roles, I have consistently delivered results in {self._format_list(self._tailor_key_skills(resume, job)[:3])}, "
                f"and I would bring that same practical, results-focused approach to this position."
            )
        if len(evidence_lines) == 1:
            return evidence_lines[0]
        return " ".join(evidence_lines[:3])

    def _select_cover_letter_evidence(self, resume: ResumeData, job: Job) -> list[str]:
        job_terms = self._job_terms(job)
        candidates: list[tuple[int, int, str]] = []
        for entry in resume.work_experience_entries:
            for idx, bullet in enumerate(entry.bullets):
                score = self._bullet_relevance_score(bullet, job_terms, job)
                sentence = self._cover_letter_sentence(entry.role_line, bullet, job_terms, job)
                candidates.append((score, -idx, sentence))
        candidates.sort(reverse=True)
        selected: list[str] = []
        seen: set[str] = set()
        for relevance, _idx, sentence in candidates:
            normalized = sentence.lower()
            if normalized in seen:
                continue
            if relevance <= 0 and selected:
                continue
            seen.add(normalized)
            selected.append(sentence)
            if len(selected) >= 3:
                break
        return selected

    def _cover_letter_sentence(self, role_line: str, bullet: str, job_terms: set[str], job: Job) -> str:
        cleaned_role = role_line.split("—", 1)[-1].strip() if "—" in role_line else role_line
        bullet_text = self._tailor_bullet_text(self._normalize_output_text(bullet), job_terms, job).rstrip(".")
        return f"As {cleaned_role}, I {self._lowercase_first_character(bullet_text)}."

    def _resume_value_summary(self, resume: ResumeData, job: Job) -> str:
        themes = self._job_theme_labels(job)
        if themes:
            return self._format_list(themes[:3]) + " that align well with the role"
        summary = self._clean_summary(resume.summary)
        if summary and summary != "an adaptable background with relevant professional experience.":
            return summary.rstrip(".").lower()
        skills = self._tailor_key_skills(resume, job)[:3]
        return self._format_list(skills) + " experience"

    def _cover_letter_closing(self, job: Job, resume: ResumeData, score: MatchScore, employer: str) -> str:
        themes = self._job_theme_labels(job)
        theme_text = self._format_list(themes[:2]) if themes else self._format_list(self._tailor_key_skills(resume, job)[:2])
        return (
            f"I would welcome the opportunity to bring my experience in {theme_text} to {employer} "
            f"and help advance the priorities of the {job.title} role."
        )

    def _infer_skills(self, resume: ResumeData, job: Job) -> list[SkillCandidate]:
        resume_text = f"{resume.raw_text} {' '.join(resume.experience_lines)} {' '.join(resume.key_skills_lines)}".lower()
        job_text = f"{job.title} {job.description_full}".lower()
        results: list[SkillCandidate] = []
        for rule in _INFERENCE_RULES:
            if not any(trigger in job_text for trigger in rule["triggers"]):
                continue
            if any(blocked in rule["skill"].lower() for blocked in _INFERENCE_BLOCKLIST):
                continue
            evidence_hits = sum(1 for evidence in rule["evidence"] if evidence in resume_text)
            if evidence_hits == 0:
                continue
            confidence = min(0.55 + (0.12 * evidence_hits), 0.92)
            if confidence > 0.50:
                results.append(SkillCandidate(rule["skill"], confidence, False))
        return results

    def _refine_bullet(self, bullet: str, job_terms: set[str]) -> str:
        cleaned = " ".join(bullet.split())
        if not cleaned:
            return cleaned
        if cleaned.endswith("."):
            return cleaned
        if len(set(re.findall(r"[a-z][a-z0-9+#&-]{2,}", cleaned.lower())) & job_terms) >= 2:
            return f"{cleaned}."
        return cleaned

    def _tailor_bullet_text(self, bullet: str, job_terms: set[str], job: Job, *, compact_level: int = 0) -> str:
        return self._tailor_bullet_text_with_mode(bullet, job_terms, job, compact_level=compact_level, allow_append=True)

    def _tailor_bullet_text_with_mode(
        self,
        bullet: str,
        job_terms: set[str],
        job: Job,
        *,
        compact_level: int = 0,
        allow_append: bool = True,
    ) -> str:
        cleaned = self._compact_bullet_text(
            self._normalize_resume_bullet_text(self._refine_bullet(bullet, job_terms)),
            compact_level,
        )
        lowered = cleaned.lower()
        themes = self._bullet_theme_labels(lowered, job_terms, job)
        if not themes:
            return cleaned
        return self._rewrite_bullet_naturally(cleaned, themes, job_terms, compact_level=compact_level, allow_append=allow_append)

    def _force_tailored_emphasis(self, bullet: str, job_terms: set[str], job: Job) -> str:
        cleaned = self._refine_bullet(bullet, job_terms).rstrip(".")
        themes = self._job_theme_labels(job)
        if not themes:
            return f"{cleaned}, aligned to the role's core priorities."
        return f"{cleaned}, emphasizing {self._format_list(themes[:2]).lower()}."

    @staticmethod
    def _skills_from_lines(lines: list[str]) -> list[str]:
        skills: list[str] = []
        for item in re.split(r"[,;\n]", " ".join(lines)):
            cleaned = item.strip(" ,")
            if cleaned:
                skills.append(cleaned)
        return skills

    @staticmethod
    def _job_terms(job: Job) -> set[str]:
        text = f"{job.title} {job.description_full}".lower()
        return {token for token in re.findall(r"[a-z][a-z0-9+#&-]{2,}", text) if token not in _STOPWORDS}

    @staticmethod
    def _bullet_relevance_score(bullet: str, job_terms: set[str], job: Job | None = None) -> int:
        words = set(re.findall(r"[a-z][a-z0-9+#&-]{2,}", bullet.lower()))
        score = len(words & job_terms)
        if job is None:
            return score
        lowered = bullet.lower()
        theme_evidence = {
            "digital growth": {"campaign", "marketing", "audience", "brand", "advertising", "acquisition", "growth"},
            "analytics and reporting": {"analytics", "reporting", "dashboard", "metrics", "analysis", "performance", "kpi", "data"},
            "cross-functional execution": {"collaborated", "coordinated", "partners", "stakeholders", "team", "cross-functional", "vendors"},
            "operational leadership": {"managed", "oversaw", "operations", "implementation", "execution", "process", "streamlined"},
            "site optimization": {"website", "optimization", "testing", "conversion", "digital", "experience"},
            "merchandising and audience strategy": {"launch", "calendar", "seasonal", "audience", "product", "promotion", "merchandising"},
            "customer acquisition and retention": {"crm", "mailchimp", "salesforce", "lead", "supporters", "customer", "retention", "email"},
            "campaign planning and launches": {"launch", "campaign", "calendar", "promotions", "creative", "go-to-market"},
            "kpi ownership": {"budget", "forecast", "revenue", "roi", "ebitda", "kpi", "performance"},
            "conversion optimization": {"conversion", "testing", "checkout", "site", "optimization", "aov"},
            "b2b marketing strategy": {"market", "segment", "customer", "strategy", "business", "industry", "growth"},
            "sales enablement and lead generation": {"lead", "pipeline", "sales", "conversion", "donor", "supporters", "outreach"},
            "lifecycle and roi reporting": {"roi", "dashboard", "reporting", "kpi", "metrics", "resource allocation", "performance"},
            "product marketing and positioning": {"positioning", "messaging", "content", "creative", "brand", "press", "white papers"},
        }
        for theme in DocumentGenerator._job_theme_labels_static(job_terms):
            score += sum(1 for token in theme_evidence.get(theme, set()) if token in lowered)
        return score

    @staticmethod
    def _skill_relevance_score(skill: str, job_terms: set[str]) -> int:
        words = set(re.findall(r"[a-z][a-z0-9+#&-]{2,}", skill.lower()))
        return len(words & job_terms)

    @staticmethod
    def _opening_stem(text: str) -> str:
        words = re.findall(r"[A-Za-z0-9+#&'-]+", text.lower())
        return " ".join(words[:4])

    @staticmethod
    def _display_employer(employer: str) -> str:
        cleaned = (employer or "").strip()
        if not cleaned or cleaned.lower() in {"nan", "none", "null"}:
            return "the hiring team"
        return cleaned

    @staticmethod
    def _clean_summary(summary: str) -> str:
        cleaned = " ".join((summary or "").split())
        lowered = cleaned.lower()
        if (
            not cleaned
            or lowered in {"work experience", "education", "key skills"}
            or re.match(r"^(january|february|march|april|may|june|july|august|september|october|november|december)\b", lowered)
            or "—" in cleaned
        ):
            return "an adaptable background with relevant professional experience."
        return cleaned.rstrip(".") + "."

    @staticmethod
    def _clean_rationale(rationale: str) -> str:
        cleaned = " ".join((rationale or "").split())
        if not cleaned or "ai scoring skipped in semi_auto" in cleaned.lower():
            return "I believe my background is relevant to the role and would welcome the opportunity to discuss it further."
        return cleaned

    @staticmethod
    def _normalize_output_text(text: str) -> str:
        normalized = " ".join((text or "").split())
        replacements = {
            "â€”": "—",
            "â€“": "–",
            "â€¢": "•",
            "â€™": "'",
            "â€œ": '"',
            "â€": '"',
        }
        for bad, good in replacements.items():
            normalized = normalized.replace(bad, good)
        return normalized

    @staticmethod
    def _normalize_resume_bullet_text(text: str) -> str:
        normalized = DocumentGenerator._normalize_output_text(text)
        replacements = {
            "crm": "CRM",
            "seo": "SEO",
            "roi": "ROI",
            "kpi": "KPI",
            "b2b": "B2B",
            "p&l": "P&L",
            "ux/ui": "UX/UI",
        }
        for bad, good in replacements.items():
            normalized = re.sub(rf"\b{re.escape(bad)}\b", good, normalized, flags=re.IGNORECASE)
        normalized = re.sub(r"\s{2,}", " ", normalized)
        return normalized.strip()

    def _signoff_name(self, resume: ResumeData) -> str:
        cleaned = self._normalize_output_text(resume.name).strip()
        if (
            not cleaned
            or cleaned.lower() in {"unknown candidate", "work experience", "education", "key skills"}
            or "@" in cleaned
            or len(cleaned.split()) < 2
        ):
            header_candidates = [
                self._normalize_output_text(line).strip()
                for line in resume.header_lines
                if line.strip()
            ]
            for candidate in header_candidates:
                if "@" in candidate or any(ch.isdigit() for ch in candidate):
                    continue
                if candidate.lower() in {"work experience", "education", "key skills"}:
                    continue
                if len(candidate.split()) >= 2:
                    return candidate.title() if candidate.isupper() else candidate
            return "Robert Thom"
        return cleaned.title() if cleaned.isupper() else cleaned

    @staticmethod
    def _lowercase_first_character(text: str) -> str:
        if not text:
            return text
        return text[0].lower() + text[1:]

    @staticmethod
    def _format_list(items: list[str]) -> str:
        cleaned = [item for item in items if item]
        if not cleaned:
            return "relevant experience"
        if len(cleaned) == 1:
            return cleaned[0]
        if len(cleaned) == 2:
            return f"{cleaned[0]} and {cleaned[1]}"
        return f"{', '.join(cleaned[:-1])}, and {cleaned[-1]}"

    def _job_theme_labels(self, job: Job) -> list[str]:
        return self._job_theme_labels_static(self._job_terms(job))

    @staticmethod
    def _job_theme_labels_static(job_terms: set[str]) -> list[str]:
        themes: list[str] = []
        if any(term in job_terms for term in {"ecommerce", "growth", "marketing", "customer", "brand"}):
            themes.append("digital growth")
        if any(term in job_terms for term in {"analytics", "reporting", "dashboard", "insights"}):
            themes.append("analytics and reporting")
        if any(term in job_terms for term in {"cross-functional", "collaboration", "stakeholder", "partnership"}):
            themes.append("cross-functional execution")
        if any(term in job_terms for term in {"operations", "process", "program", "execution"}):
            themes.append("operational leadership")
        if any(term in job_terms for term in {"site", "website", "ux", "ui", "cro", "conversion", "checkout", "seo"}):
            themes.append("site optimization")
        if any(term in job_terms for term in {"merchandising", "assortment", "pricing", "audience", "category", "product", "launch"}):
            themes.append("merchandising and audience strategy")
        if any(term in job_terms for term in {"retention", "loyalty", "repeat", "lifecycle", "acquisition", "cac", "ltv"}):
            themes.append("customer acquisition and retention")
        if any(term in job_terms for term in {"calendar", "launch", "seasonal", "campaign", "promotions"}):
            themes.append("campaign planning and launches")
        if any(term in job_terms for term in {"budget", "forecast", "revenue", "roi", "ebitda", "p&l", "profitability"}):
            themes.append("kpi ownership")
        if any(term in job_terms for term in {"conversion", "cro", "aov", "testing", "optimization"}):
            themes.append("conversion optimization")
        if any(term in job_terms for term in {"b2b", "business", "industries", "market", "markets", "segments", "segment"}):
            themes.append("b2b marketing strategy")
        if any(term in job_terms for term in {"sales", "pipeline", "lead", "leads", "closed-won", "won", "retention"}):
            themes.append("sales enablement and lead generation")
        if any(term in job_terms for term in {"roi", "lifecycle", "dashboard", "kpi", "performance", "resource", "allocation"}):
            themes.append("lifecycle and roi reporting")
        if any(term in job_terms for term in {"positioning", "value", "propositions", "product", "messaging", "case", "testimonials", "competitive"}):
            themes.append("product marketing and positioning")
        return themes

    def _bullet_theme_labels(self, bullet_text: str, job_terms: set[str], job: Job) -> list[str]:
        themes = self._job_theme_labels(job)
        labels: list[str] = []
        if "analytics and reporting" in themes and any(token in bullet_text for token in {"dashboard", "report", "metric", "analysis", "performance", "data"}):
            labels.append("analytics and reporting")
        if "digital growth" in themes and any(token in bullet_text for token in {"campaign", "advertising", "marketing", "brand", "audience", "lead"}):
            labels.append("digital growth")
        if "cross-functional execution" in themes and any(token in bullet_text for token in {"collaborated", "coordinated", "worked with", "partners", "vendors", "team"}):
            labels.append("cross-functional execution")
        if "operational leadership" in themes and any(token in bullet_text for token in {"managed", "oversaw", "streamlined", "implemented", "operations", "execution"}):
            labels.append("operational leadership")
        if "site optimization" in themes and any(token in bullet_text for token in {"website", "digital", "testing", "optimization", "performance"}):
            labels.append("site optimization")
        if "merchandising and audience strategy" in themes and any(token in bullet_text for token in {"launch", "campaign", "audience", "promotion", "creative"}):
            labels.append("merchandising and audience strategy")
        if "customer acquisition and retention" in themes and any(token in bullet_text for token in {"lead", "supporters", "crm", "mailchimp", "salesforce", "email"}):
            labels.append("customer acquisition and retention")
        if "campaign planning and launches" in themes and any(token in bullet_text for token in {"launch", "calendar", "campaign", "creative", "promotion"}):
            labels.append("campaign planning and launches")
        if "kpi ownership" in themes and any(token in bullet_text for token in {"budget", "forecast", "roi", "performance", "revenue", "metrics"}):
            labels.append("kpi ownership")
        if "conversion optimization" in themes and any(token in bullet_text for token in {"conversion", "optimization", "testing", "performance", "site"}):
            labels.append("conversion optimization")
        if "b2b marketing strategy" in themes and any(token in bullet_text for token in {"strategy", "market", "brand", "audience", "campaign", "clients"}):
            labels.append("b2b marketing strategy")
        if "sales enablement and lead generation" in themes and any(token in bullet_text for token in {"lead", "supporters", "donors", "sales", "outreach", "conversion"}):
            labels.append("sales enablement and lead generation")
        if "lifecycle and roi reporting" in themes and any(token in bullet_text for token in {"dashboard", "reporting", "metrics", "roi", "resource", "performance", "analysis"}):
            labels.append("lifecycle and roi reporting")
        if "product marketing and positioning" in themes and any(token in bullet_text for token in {"creative", "content", "press", "white papers", "brand", "messaging"}):
            labels.append("product marketing and positioning")
        return labels

    @staticmethod
    def _compact_bullet_text(text: str, compact_level: int) -> str:
        cleaned = " ".join((text or "").split())
        if compact_level <= 0:
            return cleaned
        clauses = re.split(r"(?<=[,;])\s+|\s+\band\b\s+", cleaned)
        if compact_level == 1 and len(clauses) > 1:
            cleaned = " ".join(clauses[:2]).rstrip(",;")
        if compact_level >= 2:
            cleaned = clauses[0].rstrip(",;")
            cleaned = re.sub(r"\b(to support|supporting|through|including|while|with)\b.*$", "", cleaned, flags=re.IGNORECASE).rstrip(",; ")
        return cleaned.rstrip(".") + ("." if text.endswith(".") else "")

    def _rewrite_bullet_naturally(
        self,
        bullet: str,
        themes: list[str],
        job_terms: set[str],
        *,
        compact_level: int,
        allow_append: bool = True,
    ) -> str:
        text = bullet.rstrip(".")
        words = set(re.findall(r"[a-z][a-z0-9+#&-]{2,}", text.lower()))
        overlap = len(words & job_terms)
        if overlap >= (3 if compact_level <= 1 else 2):
            return text + "."

        replacements = [
            (r"\bsocial media campaigns\b", "integrated campaigns"),
            (r"\blead generation targets\b", "lead generation and performance targets"),
            (r"\bpress releases\b", "press releases and campaign messaging"),
            (r"\bportfolio of over\b", "portfolio of"),
            (r"\bworked with\b", "partnered with"),
            (r"\bco-ordinated\b", "Coordinated"),
            (r"\bcollect data\b", "Collected data"),
            (r"\bthat report\b", "Presented that report"),
        ]
        for pattern, replacement in replacements:
            text = re.sub(pattern, replacement, text, flags=re.IGNORECASE)

        if self._bullet_has_strong_theme_evidence(text, themes):
            return text.rstrip(".") + "."

        if not allow_append:
            return text.rstrip(".") + "."

        qualifier = self._qualifier_for_themes(themes, compact_level=compact_level)
        if qualifier and not self._text_already_covers_qualifier(text, qualifier):
            text = f"{text} {qualifier}"
        return text.rstrip(".") + "."

    @staticmethod
    def _qualifier_for_themes(themes: list[str], *, compact_level: int) -> str:
        theme_priority = [
            "lifecycle and roi reporting",
            "analytics and reporting",
            "sales enablement and lead generation",
            "product marketing and positioning",
            "b2b marketing strategy",
            "digital growth",
            "cross-functional execution",
            "operational leadership",
            "site optimization",
            "merchandising and audience strategy",
            "customer acquisition and retention",
            "campaign planning and launches",
            "kpi ownership",
            "conversion optimization",
        ]
        qualifier_map = {
            "analytics and reporting": "to sharpen KPI visibility and performance decisions",
            "digital growth": "to support growth marketing and audience development",
            "cross-functional execution": "while coordinating execution across teams",
            "operational leadership": "with disciplined operational follow-through",
            "site optimization": "to strengthen digital experience performance",
            "merchandising and audience strategy": "to align campaigns with audience strategy",
            "customer acquisition and retention": "to improve acquisition and retention performance",
            "campaign planning and launches": "to support campaign planning and launch execution",
            "kpi ownership": "with clear ROI and performance accountability",
            "conversion optimization": "to improve conversion performance",
            "b2b marketing strategy": "in support of B2B growth strategy",
            "sales enablement and lead generation": "to improve lead quality and sales alignment",
            "lifecycle and roi reporting": "with a focus on lifecycle and ROI reporting",
            "product marketing and positioning": "to sharpen positioning and market-facing messaging",
        }
        compact_qualifier_map = {
            "analytics and reporting": "for KPI visibility",
            "digital growth": "for growth marketing",
            "cross-functional execution": "across teams",
            "operational leadership": "with operational discipline",
            "site optimization": "for digital performance",
            "merchandising and audience strategy": "for audience alignment",
            "customer acquisition and retention": "for acquisition and retention",
            "campaign planning and launches": "for launch execution",
            "kpi ownership": "with ROI accountability",
            "conversion optimization": "for conversion gains",
            "b2b marketing strategy": "for B2B growth",
            "sales enablement and lead generation": "for sales alignment",
            "lifecycle and roi reporting": "for lifecycle ROI visibility",
            "product marketing and positioning": "for stronger positioning",
        }
        source = compact_qualifier_map if compact_level >= 2 else qualifier_map
        ordered_themes = sorted(themes, key=lambda theme: theme_priority.index(theme) if theme in theme_priority else len(theme_priority))
        for theme in ordered_themes:
            qualifier = source.get(theme)
            if qualifier:
                return qualifier
        return ""

    @staticmethod
    def _text_already_covers_qualifier(text: str, qualifier: str) -> bool:
        text_words = set(re.findall(r"[a-z][a-z0-9+#&-]{2,}", text.lower()))
        qualifier_words = set(re.findall(r"[a-z][a-z0-9+#&-]{2,}", qualifier.lower()))
        return len(text_words & qualifier_words) >= 2

    @staticmethod
    def _bullet_has_strong_theme_evidence(text: str, themes: list[str]) -> bool:
        lowered = text.lower()
        evidence = {
            "analytics and reporting": {"dashboard", "reporting", "metrics", "analysis", "performance", "data"},
            "digital growth": {"campaign", "marketing", "brand", "audience", "growth", "engagement"},
            "cross-functional execution": {"collaborated", "coordinated", "partners", "team", "vendors"},
            "operational leadership": {"managed", "oversaw", "implemented", "operations", "execution"},
            "sales enablement and lead generation": {"lead", "sales", "supporters", "outreach", "conversion"},
            "lifecycle and roi reporting": {"roi", "dashboard", "reporting", "kpi", "metrics", "performance"},
            "product marketing and positioning": {"messaging", "creative", "content", "press", "white papers", "brand"},
        }
        for theme in themes:
            theme_evidence = evidence.get(theme, set())
            if sum(1 for token in theme_evidence if token in lowered) >= 2:
                return True
        return False

    def _build_fit_attempts(self, resume: ResumeData, job: Job) -> list[ResumeFitOptions]:
        attempts: list[ResumeFitOptions] = [
            ResumeFitOptions(compact_level=0, phase="initial"),
            ResumeFitOptions(compact_level=1, phase="tighten"),
            ResumeFitOptions(compact_level=2, phase="tighten"),
        ]
        max_global_trims = sum(max(0, len(entry.bullets) - 1) for entry in resume.work_experience_entries)
        attempts.append(ResumeFitOptions(compact_level=2, per_role_trim_rounds=1, phase="trim_bullets"))
        for global_trim_count in range(1, max_global_trims + 1):
            attempts.append(
                ResumeFitOptions(
                    compact_level=2,
                    per_role_trim_rounds=1,
                    global_trim_count=global_trim_count,
                    phase="trim_bullets",
                )
            )
        skills = self._tailor_key_skills(resume, job)
        min_skills = min(6, len(skills)) if skills else 0
        if skills and len(skills) > min_skills:
            for limit in range(len(skills) - 1, min_skills - 1, -1):
                attempts.append(
                    ResumeFitOptions(
                        compact_level=2,
                        per_role_trim_rounds=1,
                        global_trim_count=max_global_trims,
                        key_skills_limit=limit,
                        phase="shorten_skills",
                    )
                )
        return attempts

    @staticmethod
    def _fit_attempt_message(fit_options: ResumeFitOptions) -> str:
        if fit_options.phase == "tighten":
            return "Resume is over 2 pages; tightening bullets..."
        if fit_options.phase == "trim_bullets":
            return "Resume is over 2 pages; trimming lower-priority bullets..."
        if fit_options.phase == "shorten_skills":
            return "Resume is over 2 pages; shortening key skills..."
        return "Tailoring DOCX resume..."

    @staticmethod
    def _promote_staged_file(source: Path, destination: Path) -> Path:
        if source == Path() or not source or not source.exists():
            return Path()
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, destination)
        try:
            source.unlink()
        except OSError:
            pass
        return destination

    def _write_fallback_pdf(self, path: Path, job: Job, resume: ResumeData, score: MatchScore) -> None:
        lines = [
            "%PDF-1.1",
            "1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj",
            "2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj",
            "3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj",
        ]
        text_lines = [
            resume.name,
            resume.email,
            resume.phone,
            f"Target: {job.title} at {job.employer}",
            resume.summary,
            f"Match score: {score.score if score.score is not None else 'N/A'}",
            score.rationale or "No rationale available.",
        ] + resume.experience_lines[:5]
        content_stream = ["BT /F1 12 Tf 50 750 Td"]
        for idx, text in enumerate(text_lines):
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            content_stream.append(f"0 -{16 if idx else 0} Td ({escaped}) Tj")
        content_stream.append("ET")
        stream = "\n".join(content_stream)
        lines.append(f"4 0 obj << /Length {len(stream)} >> stream\n{stream}\nendstream endobj")
        lines.append("5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj")
        offsets = []
        output = []
        for item in lines:
            offsets.append(sum(len(part.encode('utf-8')) for part in output))
            output.append(f"{item}\n")
        xref_start = sum(len(part.encode("utf-8")) for part in output)
        xref = ["xref", "0 6", "0000000000 65535 f "]
        for offset in offsets:
            xref.append(f"{offset:010d} 00000 n ")
        trailer = ["trailer << /Root 1 0 R /Size 6 >>", f"startxref\n{xref_start}", "%%EOF"]
        path.write_text("".join(output + [line + "\n" for line in xref + trailer]), encoding="latin-1")
