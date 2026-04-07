from __future__ import annotations

import json
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

try:
    import pdfplumber
except ImportError:  # pragma: no cover
    pdfplumber = None

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover
    DocxDocument = None


LOGGER = logging.getLogger(__name__)
SECTION_HEADERS = {
    "summary",
    "professional summary",
    "profile",
    "objective",
    "skills",
    "key skills",
    "technical skills",
    "work experience",
    "experience",
    "professional experience",
    "education",
}
MONTH_PATTERN = re.compile(
    r"^(?:JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST|SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)"
    r"(?:\s+\d{4})?(?:\s*-\s*(?:PRESENT|(?:JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST|SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\s+\d{4}))?$",
    re.IGNORECASE,
)
MOJIBAKE_REPLACEMENTS = {
    "â€”": "—",
    "â€“": "–",
    "â€¢": "•",
    "â—": "●",
    "â€™": "'",
    "â€œ": '"',
    "â€": '"',
    "Â ": " ",
    "\uf0b7": "•",
}


@dataclass(slots=True)
class ResumeParagraphTemplate:
    paragraph_index: int
    style_name: str
    has_numbering: bool
    left_indent: int | None = None
    first_line_indent: int | None = None

    def to_dict(self) -> dict:
        return {
            "paragraph_index": self.paragraph_index,
            "style_name": self.style_name,
            "has_numbering": self.has_numbering,
            "left_indent": self.left_indent,
            "first_line_indent": self.first_line_indent,
        }


@dataclass(slots=True)
class ResumeWorkEntry:
    role_line: str
    date_line: str
    bullets: list[str] = field(default_factory=list)
    role_template: ResumeParagraphTemplate | None = None
    date_template: ResumeParagraphTemplate | None = None
    bullet_templates: list[ResumeParagraphTemplate] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "role_line": self.role_line,
            "date_line": self.date_line,
            "bullets": self.bullets,
            "role_template": self.role_template.to_dict() if self.role_template else None,
            "date_template": self.date_template.to_dict() if self.date_template else None,
            "bullet_templates": [template.to_dict() for template in self.bullet_templates],
        }


@dataclass(slots=True)
class ResumeData:
    source_path: str
    raw_text: str
    name: str
    email: str
    phone: str
    summary: str
    skills: list[str]
    experience_lines: list[str]
    header_lines: list[str] = field(default_factory=list)
    education_lines: list[str] = field(default_factory=list)
    key_skills_lines: list[str] = field(default_factory=list)
    work_experience_entries: list[ResumeWorkEntry] = field(default_factory=list)
    key_skills_templates: list[ResumeParagraphTemplate] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "source_path": self.source_path,
            "raw_text": self.raw_text,
            "name": self.name,
            "email": self.email,
            "phone": self.phone,
            "summary": self.summary,
            "skills": self.skills,
            "experience_lines": self.experience_lines,
            "header_lines": self.header_lines,
            "education_lines": self.education_lines,
            "key_skills_lines": self.key_skills_lines,
            "work_experience_entries": [entry.to_dict() for entry in self.work_experience_entries],
            "key_skills_templates": [template.to_dict() for template in self.key_skills_templates],
        }


def _normalize_text(text: str) -> str:
    normalized = text
    for bad, good in MOJIBAKE_REPLACEMENTS.items():
        normalized = normalized.replace(bad, good)
    return normalized


def _extract_text_from_path(path: Path) -> str:
    if path.suffix.lower() == ".pdf":
        if pdfplumber is None:
            raise RuntimeError("pdfplumber is required to parse PDF resumes")
        pages: list[str] = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
        return _normalize_text("\n".join(pages).strip())
    if path.suffix.lower() == ".docx":
        if DocxDocument is None:
            raise RuntimeError("python-docx is required to parse DOCX resumes")
        document = DocxDocument(str(path))
        paragraphs: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
            if "bullet" in style_name and not text.startswith(("•", "-", "*")):
                text = f"• {text}"
            paragraphs.append(text)
        return _normalize_text("\n".join(paragraphs).strip())
    return _normalize_text(path.read_text(encoding="utf-8").strip())


def _paragraph_has_numbering(paragraph) -> bool:
    p_pr = getattr(paragraph._p, "pPr", None)
    return bool(p_pr is not None and getattr(p_pr, "numPr", None) is not None)


def _paragraph_template(paragraph, paragraph_index: int) -> ResumeParagraphTemplate:
    left_indent = paragraph.paragraph_format.left_indent
    first_line_indent = paragraph.paragraph_format.first_line_indent
    return ResumeParagraphTemplate(
        paragraph_index=paragraph_index,
        style_name=paragraph.style.name if paragraph.style else "",
        has_numbering=_paragraph_has_numbering(paragraph),
        left_indent=int(left_indent) if left_indent is not None else None,
        first_line_indent=int(first_line_indent) if first_line_indent is not None else None,
    )


def _find_line(lines: list[str], token: str) -> str:
    token = token.lower()
    for line in lines:
        if token in line.lower():
            return line.strip()
    return ""


def _extract_email(text: str) -> str:
    match = re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", text)
    return match.group(0) if match else ""


def _extract_phone(text: str) -> str:
    match = re.search(r"(?:\+?1[\s.-]?)?(?:\(?\d{3}\)?[\s.-]?)\d{3}[\s.-]?\d{4}", text)
    return match.group(0) if match else ""


def _normalize_header(line: str) -> str:
    return re.sub(r"[:\s]+$", "", _normalize_text(line).strip().lower())


def _is_section_header(line: str) -> bool:
    normalized = _normalize_header(line)
    return normalized in SECTION_HEADERS


def _is_contact_or_link_line(line: str) -> bool:
    lowered = line.lower()
    return bool(
        _extract_email(line)
        or _extract_phone(line)
        or "linkedin.com" in lowered
        or "http://" in lowered
        or "https://" in lowered
        or "|" in line
    )


def _is_noise_line(line: str) -> bool:
    lowered = line.lower()
    return _is_contact_or_link_line(line) or lowered in {"work experience", "experience", "key skills", "skills"}


def _summary_from_sections(lines: list[str]) -> str:
    for idx, line in enumerate(lines):
        if _normalize_header(line) in {"summary", "professional summary", "profile", "objective"}:
            collected: list[str] = []
            for candidate in lines[idx + 1:]:
                if _is_section_header(candidate):
                    break
                if _is_noise_line(candidate):
                    continue
                if _is_role_line(candidate) or _is_date_line(candidate):
                    continue
                collected.append(_normalize_text(candidate))
                if len(" ".join(collected)) >= 240:
                    break
            if collected:
                return " ".join(collected)[:280]
    return ""


def _fallback_summary(lines: list[str]) -> str:
    candidates = [
        _normalize_text(line) for line in lines[:12]
        if not _is_noise_line(line) and not _is_section_header(line) and len(line.split()) >= 4 and not _is_role_line(line) and not _is_date_line(line)
    ]
    return " ".join(candidates[:2])[:280] if candidates else ""


def _extract_skills(lines: list[str]) -> list[str]:
    skills: list[str] = []
    for idx, line in enumerate(lines):
        normalized = _normalize_header(line)
        if normalized in {"skills", "key skills", "technical skills"} or any(normalized.startswith(f"{header}:") for header in ("skills", "key skills", "technical skills")):
            block: list[str] = []
            for candidate in lines[idx + 1:]:
                if _is_section_header(candidate):
                    break
                if _is_contact_or_link_line(candidate):
                    continue
                block.append(candidate)
                if len(block) >= 5:
                    break
            source = [line.split(":", 1)[-1]] if ":" in line else []
            source.extend(block)
            for item in re.split(r"[,;\n]", " ".join(source)):
                cleaned = item.strip(" -*\t")
                if cleaned and len(cleaned) > 1:
                    skills.append(cleaned)
            break
    deduped: list[str] = []
    seen: set[str] = set()
    for skill in skills:
        lowered = skill.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        deduped.append(skill)
    return deduped[:12]


def _extract_experience_lines(lines: list[str]) -> list[str]:
    experience: list[str] = []
    in_experience_section = False
    for line in lines:
        normalized = _normalize_header(line)
        if normalized in {"work experience", "experience", "professional experience"}:
            in_experience_section = True
            continue
        if in_experience_section and _is_section_header(line):
            break
        if _is_noise_line(line):
            continue
        if in_experience_section:
            if any(keyword in line.lower() for keyword in ("manager", "director", "analyst", "coordinator", "specialist", "developer", "engineer", "campaign", "operations")):
                experience.append(line)
        elif any(keyword in line.lower() for keyword in ("engineer", "developer", "manager", "analyst", "director", "coordinator", "specialist", "built", "designed")):
            experience.append(line)
    deduped: list[str] = []
    seen: set[str] = set()
    for line in experience:
        lowered = line.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        deduped.append(line)
    return deduped[:8]


def _find_section_index(lines: list[str], names: tuple[str, ...]) -> int:
    for idx, line in enumerate(lines):
        if _normalize_header(line) in names:
            return idx
    return -1


def _is_role_line(line: str) -> bool:
    return "—" in line or " - " in line or " – " in line


def _is_date_line(line: str) -> bool:
    return bool(MONTH_PATTERN.match(line.strip()))


def _is_bullet_paragraph(paragraph) -> bool:
    text = paragraph.text.strip()
    if not text:
        return False
    style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
    if "bullet" in style_name:
        return True
    if _paragraph_has_numbering(paragraph):
        return True
    first_line_indent = paragraph.paragraph_format.first_line_indent
    left_indent = paragraph.paragraph_format.left_indent
    if first_line_indent is not None and left_indent is not None and int(first_line_indent) < 0 and int(left_indent) > 0:
        return True
    return text.startswith(("•", "-", "*", "●"))


def _structured_sections_from_docx(source_path: Path) -> tuple[list[str], list[ResumeWorkEntry], list[str], list[str], list[ResumeParagraphTemplate]]:
    if DocxDocument is None:  # pragma: no cover
        raise RuntimeError("python-docx is required to parse DOCX resumes")
    document = DocxDocument(str(source_path))
    paragraphs = document.paragraphs
    nonempty = [(idx, paragraph) for idx, paragraph in enumerate(paragraphs) if paragraph.text.strip()]

    def _find_heading_index_docx(names: tuple[str, ...]) -> int:
        for idx, paragraph in nonempty:
            if _normalize_header(paragraph.text) in names:
                return idx
        return -1

    work_idx = _find_heading_index_docx(("work experience", "experience", "professional experience"))
    education_idx = _find_heading_index_docx(("education",))
    skills_idx = _find_heading_index_docx(("key skills", "skills", "technical skills"))

    header_lines = [paragraph.text.strip() for idx, paragraph in nonempty if idx < work_idx] if work_idx > 0 else [paragraph.text.strip() for _, paragraph in nonempty[:2]]

    work_entries: list[ResumeWorkEntry] = []
    current: ResumeWorkEntry | None = None
    for idx, paragraph in nonempty:
        if work_idx < 0 or idx <= work_idx or (education_idx >= 0 and idx >= education_idx):
            continue
        text = paragraph.text.strip()
        if not text or text.lower() in {line.lower() for line in header_lines} or _is_contact_or_link_line(text):
            continue
        if current and not current.date_line and _is_date_line(text):
            current.date_line = text
            current.date_template = _paragraph_template(paragraph, idx)
            continue
        if _is_bullet_paragraph(paragraph):
            bullet_text = text.lstrip("•●*- ").strip()
            if current is None:
                continue
            current.bullets.append(bullet_text)
            current.bullet_templates.append(_paragraph_template(paragraph, idx))
            continue
        if _is_role_line(text):
            if current and current.role_line:
                work_entries.append(current)
            current = ResumeWorkEntry(
                role_line=text,
                date_line="",
                bullets=[],
                role_template=_paragraph_template(paragraph, idx),
                date_template=None,
                bullet_templates=[],
            )
            continue
        if current and current.bullets:
            current.bullets[-1] = f"{current.bullets[-1]} {text}".strip()
        elif current is None:
            continue
    if current and current.role_line:
        work_entries.append(current)

    education_lines = [paragraph.text.strip() for idx, paragraph in nonempty if education_idx >= 0 and idx > education_idx and (skills_idx < 0 or idx < skills_idx)]
    key_skill_pairs = [(idx, paragraph) for idx, paragraph in nonempty if skills_idx >= 0 and idx > skills_idx]
    key_skills_lines = [paragraph.text.strip() for idx, paragraph in key_skill_pairs]
    key_skills_templates = [_paragraph_template(paragraph, idx) for idx, paragraph in key_skill_pairs]
    return header_lines, work_entries, education_lines, key_skills_lines, key_skills_templates


def _parse_work_entries(lines: list[str], *, header_lines: list[str]) -> list[ResumeWorkEntry]:
    entries: list[ResumeWorkEntry] = []
    current: ResumeWorkEntry | None = None
    header_values = {line.strip().lower() for line in header_lines}
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        if line.lower() in header_values:
            continue
        if _is_contact_or_link_line(line):
            continue
        if current and not current.date_line and _is_date_line(line):
            current.date_line = line
            continue
        if _is_role_line(line):
            if current and current.role_line:
                entries.append(current)
            current = ResumeWorkEntry(role_line=line, date_line="", bullets=[])
            continue
        bullet_text = line.lstrip("●•*- ").strip()
        if current is None:
            continue
        if line.startswith(("●", "•", "-", "*")) or not current.bullets:
            current.bullets.append(bullet_text)
        else:
            current.bullets[-1] = f"{current.bullets[-1]} {bullet_text}".strip()
    if current and current.role_line:
        entries.append(current)
    return entries


def _structured_sections(lines: list[str]) -> tuple[list[str], list[ResumeWorkEntry], list[str], list[str]]:
    work_idx = _find_section_index(lines, ("work experience", "experience", "professional experience"))
    education_idx = _find_section_index(lines, ("education",))
    skills_idx = _find_section_index(lines, ("key skills", "skills", "technical skills"))
    header_lines = lines[:work_idx] if work_idx > 0 else lines[:2]
    work_lines = lines[work_idx + 1:education_idx] if work_idx >= 0 and education_idx > work_idx else []
    education_lines = lines[education_idx + 1:skills_idx] if education_idx >= 0 and skills_idx > education_idx else []
    key_skills_lines = lines[skills_idx + 1:] if skills_idx >= 0 else []
    work_entries = _parse_work_entries(work_lines, header_lines=header_lines)
    return header_lines, work_entries, education_lines, key_skills_lines


def _guess_name(lines: list[str], source_path: Path | None = None) -> str:
    for line in lines[:5]:
        line = _normalize_text(line).strip()
        if not line:
            continue
        if _is_section_header(line) or _is_role_line(line) or _is_date_line(line):
            continue
        if "@" in line or any(ch.isdigit() for ch in line):
            continue
        if len(line.split()) < 2:
            continue
        return line
    if source_path is not None:
        stem = source_path.stem
        stem = re.sub(r"\bresume\b", "", stem, flags=re.IGNORECASE)
        stem = re.sub(r"\b\d{4}\b", "", stem)
        stem = re.sub(r"[_-]+", " ", stem)
        stem = " ".join(part for part in stem.split() if part)
        if len(stem.split()) >= 2:
            return stem.title()
    return "Unknown Candidate"


def parse_resume(source_path: Path, cache_path: Path) -> ResumeData:
    text = _extract_text_from_path(source_path)
    if not text:
        raise ValueError("Resume file contained no extractable text")

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if source_path.suffix.lower() == ".docx":
        header_lines, work_entries, education_lines, key_skills_lines, key_skills_templates = _structured_sections_from_docx(source_path)
    else:
        header_lines, work_entries, education_lines, key_skills_lines = _structured_sections(lines)
        key_skills_templates = []
    email = _extract_email(text)
    phone = _extract_phone(text)
    summary_line = _summary_from_sections(lines) or _fallback_summary(lines) or "Experienced professional with relevant background."
    experience_lines = _extract_experience_lines(lines)
    skills = _extract_skills(lines)
    resume = ResumeData(
        source_path=str(source_path),
        raw_text=text,
        name=_guess_name(lines, source_path),
        email=email,
        phone=phone,
        summary=summary_line,
        skills=skills,
        experience_lines=experience_lines,
        header_lines=header_lines,
        education_lines=education_lines,
        key_skills_lines=key_skills_lines,
        work_experience_entries=work_entries,
        key_skills_templates=key_skills_templates,
    )
    cache_path.write_text(json.dumps(resume.to_dict(), indent=2), encoding="utf-8")
    LOGGER.info("Cached parsed resume at %s", cache_path)
    return resume


def load_cached_resume(cache_path: Path) -> ResumeData | None:
    if not cache_path.exists():
        return None
    payload = json.loads(cache_path.read_text(encoding="utf-8"))
    payload["header_lines"] = payload.get("header_lines", [])
    payload["education_lines"] = payload.get("education_lines", [])
    payload["key_skills_lines"] = payload.get("key_skills_lines", [])
    payload["key_skills_templates"] = [
        ResumeParagraphTemplate(**entry) if isinstance(entry, dict) else entry
        for entry in payload.get("key_skills_templates", [])
    ]
    payload["work_experience_entries"] = [
        ResumeWorkEntry(
            role_line=entry["role_line"],
            date_line=entry.get("date_line", ""),
            bullets=entry.get("bullets", []),
            role_template=ResumeParagraphTemplate(**entry["role_template"]) if isinstance(entry.get("role_template"), dict) else entry.get("role_template"),
            date_template=ResumeParagraphTemplate(**entry["date_template"]) if isinstance(entry.get("date_template"), dict) else entry.get("date_template"),
            bullet_templates=[
                ResumeParagraphTemplate(**template) if isinstance(template, dict) else template
                for template in entry.get("bullet_templates", [])
            ],
        ) if isinstance(entry, dict) else entry
        for entry in payload.get("work_experience_entries", [])
    ]
    return ResumeData(**payload)
