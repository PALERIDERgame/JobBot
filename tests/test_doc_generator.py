from __future__ import annotations

import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from database import Job
from doc_generator import DocumentGenerator
from document_tailoring import DocumentTailoringPayload
from match_scorer import MatchScore
from resume_parser import ResumeData, ResumeWorkEntry, parse_resume
from test_support import workspace_temp_dir


def _build_source_resume_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("ALEX RIVERS")
    doc.add_paragraph("alex@example.com | linkedin.com/in/example | (555) 000-1234")
    doc.add_paragraph("WORK EXPERIENCE", style="Heading 1")
    doc.add_paragraph("Underdog Strategies, New York, NY — Digital Advertising and Field Manager", style="Heading 2")
    doc.add_paragraph("JULY 2024 - Present")
    bullet = doc.add_paragraph(style="List Bullet")
    bullet.add_run("Led digital ad strategy and reporting")
    bullet = doc.add_paragraph(style="List Bullet")
    bullet.add_run("Managed creative vendors and launches")
    doc.add_paragraph("Behavioral Associates, New York, NY — Operations Manager", style="Heading 2")
    doc.add_paragraph("JUNE 2023 - JULY 2024", style="Heading 3")
    bullet = doc.add_paragraph(style="List Bullet")
    bullet.add_run("Directed software implementation and internal systems")
    doc.add_paragraph("EDUCATION", style="Heading 1")
    doc.add_paragraph("University of California, Davis", style="Heading 2")
    doc.add_paragraph("KEY SKILLS", style="Heading 1")
    doc.add_paragraph("Operations, Project Management, CRM, Data Analysis, Mailchimp, Google Analytics")
    doc.save(str(path))


class DocumentGeneratorTests(unittest.TestCase):
    def test_repair_ai_tailoring_payload_falls_back_for_bad_cover_letter_and_broken_bullets(self) -> None:
        generator = DocumentGenerator()
        local_payload = DocumentTailoringPayload(
            work_entries=[
                ResumeWorkEntry("Role 1", "2024", ["Managed and optimized social media campaigns.", "Presented the report to leadership."])
            ],
            key_skills=["Dashboard Reporting"],
            cover_letter_text="Dear Hiring Team at Acme,\n\nI am excited to apply for the Head of Marketing role.\n\nSincerely,\nJane",
            route="local",
        )
        ai_payload = DocumentTailoringPayload(
            work_entries=[
                ResumeWorkEntry("Role 1", "2024", ["Managed optimized social media campaigns.", "Presented Presented that report."])
            ],
            key_skills=["Dashboard Reporting", "Lead Generation"],
            cover_letter_text="{'type': 'string'}",
            route="openai",
            ai_attempted=True,
            provider="openai",
            model="gpt-5-mini",
        )
        repaired, issues = generator.repair_ai_tailoring_payload(ai_payload, local_payload)
        self.assertEqual(repaired.cover_letter_text, local_payload.cover_letter_text)
        self.assertEqual(repaired.cover_letter_ai_status, "local")
        self.assertEqual(repaired.cover_letter_fallback, "local")
        self.assertEqual(repaired.resume_ai_status, "local")
        self.assertEqual(repaired.rejected_bullets_repaired, 2)
        self.assertEqual(repaired.work_entries[0].bullets, local_payload.work_entries[0].bullets)
        self.assertTrue(any(issue.startswith("cover_letter_") for issue in issues))
        self.assertTrue(any("broken_conjunction" in issue or "duplicate_word" in issue for issue in issues))

    def test_repair_ai_tailoring_payload_keeps_valid_bullets_and_marks_partial(self) -> None:
        generator = DocumentGenerator()
        local_payload = DocumentTailoringPayload(
            work_entries=[
                ResumeWorkEntry("Role 1", "2024", ["Managed and optimized social media campaigns.", "Presented the report to leadership."])
            ],
            key_skills=["Dashboard Reporting"],
            cover_letter_text="Dear Hiring Team at Acme,\n\nI am excited to apply for the Head of Marketing role with a background in reporting and operations.\n\nSincerely,\nJane",
            route="local",
        )
        ai_payload = DocumentTailoringPayload(
            work_entries=[
                ResumeWorkEntry("Role 1", "2024", ["Improved reporting cadence and dashboard visibility for leadership.", "Presented Presented that report."])
            ],
            key_skills=["Dashboard Reporting", "Lead Generation"],
            cover_letter_text="Dear Hiring Team at Acme,\n\nI am excited to apply for the Head of Marketing role and would bring experience in dashboard reporting, cross-functional execution, and audience growth.\n\nSincerely,\nJane",
            route="openai",
            ai_attempted=True,
            provider="openai",
            model="gpt-5-mini",
        )
        repaired, issues = generator.repair_ai_tailoring_payload(ai_payload, local_payload)
        self.assertEqual(repaired.route, "openai")
        self.assertEqual(repaired.resume_ai_status, "partial")
        self.assertEqual(repaired.cover_letter_ai_status, "accepted")
        self.assertEqual(repaired.rejected_bullets_repaired, 1)
        self.assertEqual(repaired.work_entries[0].bullets[0], "Improved reporting cadence and dashboard visibility for leadership.")
        self.assertEqual(repaired.work_entries[0].bullets[1], "Presented the report to leadership.")
        self.assertTrue(any("duplicate_word" in issue for issue in issues))

    def test_generate_creates_files(self) -> None:
        with workspace_temp_dir() as tmp:
            generator = DocumentGenerator()
            def _fake_export(docx_path, pdf_path):
                pdf_path.write_bytes(b"pdf")
                return "word_com"
            generator._export_docx_to_pdf = _fake_export
            docs = generator.generate(
                Path(tmp),
                Job("1", "Engineer", "Acme", "Remote", "", "Role", "board", "https://example.com", "", "usajobs", "", ""),
                ResumeData("", "", "Jane", "jane@example.com", "555", "Summary", ["Python"], ["Built APIs"]),
                MatchScore(90, "Great fit", ["Python"], ["AWS"], True, "scored", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="Strong API background and automation experience.",
            )
            self.assertTrue(docs.resume_pdf_path.exists())
            self.assertTrue(docs.cover_letter_txt_path.exists())
            self.assertTrue(docs.cover_letter_docx_path.exists())
            self.assertTrue(docs.cover_letter_pdf_path.exists())

    def test_generate_cover_letter_avoids_nan_employer_and_placeholder_rationale(self) -> None:
        with workspace_temp_dir() as tmp:
            generator = DocumentGenerator()
            def _fake_export(docx_path, pdf_path):
                pdf_path.write_bytes(b"pdf")
                return "word_com"
            generator._export_docx_to_pdf = _fake_export
            docs = generator.generate(
                Path(tmp),
                Job("1", "Marketing Director", "nan", "New York, NY", "", "Role", "board", "https://example.com", "", "jobspy", "", ""),
                ResumeData("", "", "Jane", "jane@example.com", "555", "Marketing leader with operations experience", ["Campaign Management", "Operations"], ["Led field operations", "Managed paid media"]),
                MatchScore(0, "AI scoring skipped in semi_auto for faster review queueing.", [], [], False, "review", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="",
            )
            text = docs.cover_letter_txt_path.read_text(encoding="utf-8")
            self.assertNotIn("nan", text.lower())
            self.assertNotIn("AI scoring skipped", text)
            self.assertNotIn("Selected experience highlights:", text)
            self.assertNotIn("I believe my background is relevant", text)

    def test_generate_docx_resume_preserves_sections_and_writes_docx_and_pdf(self) -> None:
        with workspace_temp_dir() as tmp:
            source_docx = Path(tmp) / "source_resume.docx"
            _build_source_resume_docx(source_docx)
            generator = DocumentGenerator()
            generator._export_docx_to_pdf = lambda docx_path, pdf_path: pdf_path.write_bytes(b"pdf")
            generator._count_pdf_pages = lambda pdf_path: 2
            resume = parse_resume(source_docx, Path(tmp) / "resume_cache.json")
            docs = generator.generate(
                Path(tmp),
                Job("1", "Director of Ecommerce", "MILK BAR", "New York, NY", "", "Lead ecommerce growth, analytics, and dashboard reporting.", "board", "https://example.com", "", "jobspy", "", ""),
                resume,
                MatchScore(0, "", [], [], False, "review", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="",
            )
            self.assertTrue(docs.resume_docx_path.exists())
            self.assertTrue(docs.resume_pdf_path.exists())
            self.assertEqual(docs.status, "generated")

            from docx import Document
            document = Document(str(docs.resume_docx_path))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("WORK EXPERIENCE", text)
            self.assertIn("EDUCATION", text)
            self.assertIn("KEY SKILLS", text)
            self.assertIn("University of California, Davis", text)
            self.assertNotIn("Target Role", text)
            self.assertNotIn("Match Notes", text)
            work_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
            numbered = [paragraph for paragraph in work_paragraphs if paragraph.text.startswith("Led") or paragraph.text.startswith("Managed") or paragraph.text.startswith("Directed")]
            self.assertTrue(numbered)
            self.assertTrue(all(paragraph.style.name != "Heading 2" for paragraph in numbered))
            self.assertTrue(
                all(
                    getattr(paragraph._p.pPr, "numPr", None) is not None or "bullet" in paragraph.style.name.lower()
                    for paragraph in numbered
                )
            )
            source_document = Document(str(source_docx))
            source_bullet = next(paragraph for paragraph in source_document.paragraphs if paragraph.text.startswith("Led digital ad strategy"))
            generated_bullet = next(paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Drove") or paragraph.text.startswith("Led digital") or paragraph.text.startswith("Led"))
            self.assertIsNotNone(generated_bullet.runs[0].font.size)
            if source_bullet.runs[0].font.size is not None:
                self.assertLess(generated_bullet.runs[0].font.size, source_bullet.runs[0].font.size)
            with ZipFile(docs.resume_docx_path) as archive:
                numbering_xml = archive.read("word/numbering.xml").decode("utf-8", errors="ignore")
            self.assertIn('w:sz w:val="12"', numbering_xml)
            self.assertIn('w:szCs w:val="12"', numbering_xml)
            key_skills_heading_idx = next(idx for idx, paragraph in enumerate(work_paragraphs) if paragraph.text.strip() == "KEY SKILLS")
            self.assertEqual(len([paragraph for paragraph in work_paragraphs[key_skills_heading_idx + 1:] if paragraph.text.strip()]), 1)

    def test_generate_docx_resume_tailors_work_bullets_for_relevant_job(self) -> None:
        with workspace_temp_dir() as tmp:
            source_docx = Path(tmp) / "source_resume.docx"
            _build_source_resume_docx(source_docx)
            generator = DocumentGenerator()
            generator._export_docx_to_pdf = lambda docx_path, pdf_path: pdf_path.write_bytes(b"pdf")
            generator._count_pdf_pages = lambda pdf_path: 2
            resume = parse_resume(source_docx, Path(tmp) / "resume_cache.json")
            docs = generator.generate(
                Path(tmp),
                Job("1", "Ecommerce Growth Manager", "MILK BAR", "New York, NY", "", "Lead ecommerce growth, marketing analytics, dashboard reporting, and cross-functional execution.", "board", "https://example.com", "", "jobspy", "", ""),
                resume,
                MatchScore(0, "", [], [], False, "review", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="",
            )
            from docx import Document
            document = Document(str(docs.resume_docx_path))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
            self.assertTrue(
                any(
                    phrase in text.lower()
                    for phrase in (
                        "kpi visibility",
                        "growth marketing",
                        "acquisition and retention",
                        "across teams",
                    )
                )
            )
            self.assertNotIn("with a focus on", text.lower())
            self.assertNotIn("digital growth initiatives, campaign planning, and customer acquisition efforts by", text.lower())
            self.assertNotIn("Target Role", text)

    def test_cover_letter_is_structured_and_tailored_for_relevant_job(self) -> None:
        with workspace_temp_dir() as tmp:
            source_docx = Path(tmp) / "source_resume.docx"
            _build_source_resume_docx(source_docx)
            generator = DocumentGenerator()
            generator._export_docx_to_pdf = lambda docx_path, pdf_path: pdf_path.write_bytes(b"pdf")
            generator._count_pdf_pages = lambda pdf_path: 2
            resume = parse_resume(source_docx, Path(tmp) / "resume_cache.json")
            docs = generator.generate(
                Path(tmp),
                Job("1", "Director of Ecommerce", "MILK BAR", "New York, NY", "", "Lead ecommerce growth, marketing analytics, dashboard reporting, and cross-functional execution.", "board", "https://example.com", "", "jobspy", "", ""),
                resume,
                MatchScore(0, "", [], [], False, "review", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="",
            )
            text = docs.cover_letter_txt_path.read_text(encoding="utf-8")
            self.assertIn("Director of Ecommerce", text)
            self.assertIn("MILK BAR", text)
            self.assertIn("analytics and reporting", text.lower())
            self.assertIn("digital growth", text.lower())
            self.assertIn("Sincerely,\nAlex Rivers", text)
            self.assertNotIn("Selected experience highlights:", text)
            self.assertNotIn("I believe my background is relevant", text)
            self.assertNotIn("â€”", text)
            self.assertNotIn("- Underdog Strategies", text)

    def test_docx_generation_auto_fits_resume_to_two_pages(self) -> None:
        with workspace_temp_dir() as tmp:
            source_docx = Path(tmp) / "source_resume.docx"
            _build_source_resume_docx(source_docx)
            generator = DocumentGenerator()
            generator._export_docx_to_pdf = lambda docx_path, pdf_path: pdf_path.write_bytes(b"pdf")
            page_counts = iter([3, 2])
            generator._count_pdf_pages = lambda pdf_path: next(page_counts)
            resume = parse_resume(source_docx, Path(tmp) / "resume_cache.json")
            docs = generator.generate(
                Path(tmp),
                Job("1", "Director of Ecommerce", "MILK BAR", "New York, NY", "", "Lead ecommerce growth, analytics, dashboard reporting, and site optimization.", "board", "https://example.com", "", "jobspy", "", ""),
                resume,
                MatchScore(0, "", [], [], False, "review", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="",
            )
            self.assertEqual(docs.status, "generated")
            self.assertTrue(docs.resume_docx_path.exists())
            self.assertTrue(docs.resume_pdf_path.exists())

    def test_docx_generation_fails_only_after_exhausting_fit_attempts(self) -> None:
        with workspace_temp_dir() as tmp:
            source_docx = Path(tmp) / "source_resume.docx"
            _build_source_resume_docx(source_docx)
            generator = DocumentGenerator()
            generator._export_docx_to_pdf = lambda docx_path, pdf_path: pdf_path.write_bytes(b"pdf")
            generator._count_pdf_pages = lambda pdf_path: 3
            resume = parse_resume(source_docx, Path(tmp) / "resume_cache.json")
            docs = generator.generate(
                Path(tmp),
                Job("1", "Director of Ecommerce", "MILK BAR", "New York, NY", "", "Lead ecommerce growth, analytics, dashboard reporting, and site optimization.", "board", "https://example.com", "", "jobspy", "", ""),
                resume,
                MatchScore(0, "", [], [], False, "review", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="",
            )
            self.assertEqual(docs.status, "failed")
            self.assertIn("compress the tailored resume to 2 pages", docs.error_message.lower())
            self.assertEqual(docs.resume_docx_path, Path())
            self.assertEqual(docs.resume_pdf_path, Path())

    def test_docx_generation_emits_fit_progress_messages(self) -> None:
        with workspace_temp_dir() as tmp:
            source_docx = Path(tmp) / "source_resume.docx"
            _build_source_resume_docx(source_docx)
            generator = DocumentGenerator()
            generator._export_docx_to_pdf = lambda docx_path, pdf_path: pdf_path.write_bytes(b"pdf")
            page_counts = iter([3, 2])
            generator._count_pdf_pages = lambda pdf_path: next(page_counts)
            resume = parse_resume(source_docx, Path(tmp) / "resume_cache.json")
            events: list[tuple[str, str, int]] = []
            generator.generate(
                Path(tmp),
                Job("1", "Director of Ecommerce", "MILK BAR", "New York, NY", "", "Lead ecommerce growth, analytics, dashboard reporting, and site optimization.", "board", "https://example.com", "", "jobspy", "", ""),
                resume,
                MatchScore(0, "", [], [], False, "review", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="",
                progress_callback=lambda stage, message, progress: events.append((stage, message, progress)),
            )
            messages = [message for _stage, message, _progress in events]
            self.assertIn("Validating final page count...", messages)
            self.assertIn("Resume is over 2 pages; tightening bullets...", messages)

    def test_cover_letter_uses_safe_signoff_when_resume_name_is_bad(self) -> None:
        generator = DocumentGenerator()
        resume = ResumeData(
            "resume.docx",
            "",
            "WORK EXPERIENCE",
            "alex@example.com",
            "(555) 000-1234",
            "",
            ["Operations"],
            ["Led digital ad strategy and reporting"],
            header_lines=["ALEX RIVERS", "alex@example.com | linkedin.com/in/example | (555) 000-1234"],
            work_experience_entries=[
                ResumeWorkEntry(
                    role_line="Underdog Strategies, New York, NY — Digital Advertising and Field Manager",
                    date_line="JULY 2024 - Present",
                    bullets=["Led digital ad strategy and reporting"],
                )
            ],
        )
        with workspace_temp_dir() as tmp:
            def _fake_export(docx_path, pdf_path):
                pdf_path.write_bytes(b"pdf")
                return "word_com"
            generator._export_docx_to_pdf = _fake_export
            docs = generator.generate(
                Path(tmp),
                Job("1", "Director", "Acme", "Remote", "", "Lead analytics and reporting.", "board", "https://example.com", "", "jobspy", "", ""),
                resume,
                MatchScore(0, "", [], [], False, "review", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="",
            )
            text = docs.cover_letter_txt_path.read_text(encoding="utf-8")
            self.assertIn("Sincerely,\nAlex Rivers", text)
            self.assertNotIn("Sincerely,\nWORK EXPERIENCE", text)

    def test_generate_cover_letter_outputs_docx_and_pdf(self) -> None:
        with workspace_temp_dir() as tmp:
            source_docx = Path(tmp) / "source_resume.docx"
            _build_source_resume_docx(source_docx)
            generator = DocumentGenerator()
            def _fake_export(docx_path, pdf_path):
                pdf_path.write_bytes(b"pdf")
                return "word_com"
            generator._export_docx_to_pdf = _fake_export
            generator._count_pdf_pages = lambda pdf_path: 2
            resume = parse_resume(source_docx, Path(tmp) / "resume_cache.json")
            docs = generator.generate(
                Path(tmp),
                Job("1", "Director of Ecommerce", "MILK BAR", "New York, NY", "", "Lead ecommerce growth.", "board", "https://example.com", "", "jobspy", "", ""),
                resume,
                MatchScore(0, "", [], [], False, "review", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="",
            )
            self.assertTrue(docs.cover_letter_txt_path.exists())
            self.assertTrue(docs.cover_letter_docx_path.exists())
            self.assertTrue(docs.cover_letter_pdf_path.exists())

    def test_inferred_skills_are_added_when_supported(self) -> None:
        generator = DocumentGenerator()
        resume = ResumeData(
            "resume.docx",
            "Mailchimp Salesforce Marketing Cloud streamlined campaign reporting dashboards",
            "ALEX RIVERS",
            "alex@example.com",
            "(555) 000-1234",
            "Marketing leader",
            ["Operations"],
            ["Streamlined campaign reporting dashboards for executive leadership"],
            key_skills_lines=["Operations, Project Management, CRM, Data Analysis, Mailchimp, Salesforce Marketing Cloud"],
            work_experience_entries=[],
        )
        job = Job("1", "Lifecycle Marketing Manager", "Acme", "Remote", "", "Own email marketing, dashboard reporting, and process improvement.", "board", "", "", "jobspy", "", "")
        skills = generator._tailor_key_skills(resume, job)
        self.assertIn("Email Marketing", skills)
        self.assertIn("Dashboard Reporting", skills)
        self.assertIn("Process Improvement", skills)

    def test_unsupported_credentials_are_not_inferred(self) -> None:
        generator = DocumentGenerator()
        resume = ResumeData(
            "resume.docx",
            "Planned and ran programming and managed campaigns and analytics dashboards",
            "ALEX RIVERS",
            "alex@example.com",
            "(555) 000-1234",
            "Marketing leader",
            ["Operations"],
            ["Planned and ran programming and managed campaigns and analytics dashboards"],
            key_skills_lines=["Operations, Project Management"],
            work_experience_entries=[],
        )
        job = Job("1", "Program Manager", "Acme", "Remote", "", "PMP certification required and program management preferred.", "board", "", "", "jobspy", "", "")
        skills = generator._tailor_key_skills(resume, job)
        self.assertIn("Program Management", skills)
        self.assertNotIn("PMP certification", " ".join(skills))

    def test_docx_generation_can_succeed_without_pdf_export(self) -> None:
        with workspace_temp_dir() as tmp:
            source_docx = Path(tmp) / "source_resume.docx"
            _build_source_resume_docx(source_docx)
            generator = DocumentGenerator()
            generator._export_docx_to_pdf = lambda docx_path, pdf_path: (_ for _ in ()).throw(RuntimeError("Word automation unavailable"))
            resume = parse_resume(source_docx, Path(tmp) / "resume_cache.json")
            docs = generator.generate(
                Path(tmp),
                Job("1", "Director", "Acme", "Remote", "", "Role", "board", "", "", "jobspy", "", ""),
                resume,
                MatchScore(0, "", [], [], False, "review", "", ""),
                ai_notes="",
            )
            self.assertTrue(docs.resume_docx_path.exists())
            self.assertEqual(docs.resume_pdf_path, Path())
            self.assertEqual(docs.status, "generated_docx_only")
            self.assertIn("Word automation unavailable", docs.error_message)

    def test_docx_pdf_export_uses_common_libreoffice_path_when_not_on_path(self) -> None:
        with workspace_temp_dir() as tmp:
            source_docx = Path(tmp) / "source_resume.docx"
            pdf_path = Path(tmp) / "resume.pdf"
            _build_source_resume_docx(source_docx)
            generator = DocumentGenerator()

            class _Completed:
                returncode = 0
                stdout = ""
                stderr = ""

            def _fake_run(*args, **kwargs):
                pdf_path.write_bytes(b"pdf")
                return _Completed()

            with patch.object(generator, "_libreoffice_candidates", return_value=[r"C:\Program Files\LibreOffice\program\soffice.exe"]), patch(
                "doc_generator.subprocess.run",
                side_effect=_fake_run,
            ) as run_mock, patch.dict("sys.modules", {"win32com": None, "win32com.client": None}):
                result = generator._export_docx_to_pdf(source_docx, pdf_path)

            self.assertTrue(pdf_path.exists())
            self.assertEqual(result, "libreoffice")
            self.assertIn(r"C:\Program Files\LibreOffice\program\soffice.exe", " ".join(run_mock.call_args.args[0]))

    def test_dow_jones_style_bullets_avoid_repetitive_prefixes_and_bad_grammar(self) -> None:
        with workspace_temp_dir() as tmp:
            source_docx = Path(tmp) / "source_resume.docx"
            _build_source_resume_docx(source_docx)
            generator = DocumentGenerator()
            generator._export_docx_to_pdf = lambda docx_path, pdf_path: pdf_path.write_bytes(b"pdf")
            generator._count_pdf_pages = lambda pdf_path: 2
            resume = parse_resume(source_docx, Path(tmp) / "resume_cache.json")
            resume.work_experience_entries[0].bullets = [
                "Spearheaded a new digital ad strategy for charter school clients, including sentiment analysis of ad copy, which boosted student inquiries by 35% and lowered cost per lead by 15% across key campaigns.",
                "Managed and optimized social media campaigns, leading to a 50% increase in brand awareness and a 25% increase in engagement with key demographics in the Bronx.",
                "Streamlined lead data review and updated real-time dashboards for executive leadership, developing metrics for canvasser productivity and providing immediate insights for campaign adjustments.",
                "Led and trained a team of 15+ canvassers, exceeding monthly lead generation targets by an average of 10% through a new data-informed training program focused on effective community engagement.",
            ]
            docs = generator.generate(
                Path(tmp),
                Job("1", "Head of Marketing, Industries", "Dow Jones", "New York, NY", "", "Define marketing strategy, maximize ROI, align with sales, improve lead quality, and build performance dashboards for B2B growth.", "email", "https://example.com", "talent@example.com", "jobspy", "", ""),
                resume,
                MatchScore(0, "", [], [], False, "review", "", "2026-01-01T00:00:00+00:00"),
                ai_notes="",
            )
            from docx import Document
            document = Document(str(docs.resume_docx_path))
            bullet_text = [p.text for p in document.paragraphs if p.text.strip() and p.style.name == "List Bullet"]
            combined = "\n".join(bullet_text).lower()
            self.assertNotIn("by and", combined)
            self.assertNotIn("by a portfolio", combined)
            self.assertNotIn("by met", combined)
            self.assertNotIn("by set", combined)
            self.assertNotIn("by crm", combined)
            self.assertLess(combined.count("to support growth marketing and audience development"), 3)
            self.assertTrue(
                any(
                    phrase in combined
                    for phrase in (
                        "b2b growth strategy",
                        "sales alignment",
                        "roi reporting",
                        "kpi visibility",
                    )
                )
            )

    def test_should_escalate_tailoring_for_broken_local_output(self) -> None:
        generator = DocumentGenerator()
        job = Job("1", "Head of Marketing, Industries", "Dow Jones", "New York, NY", "", "Lead B2B growth, ROI reporting, sales alignment, and product marketing.", "email", "", "", "jobspy", "", "")
        resume = ResumeData("", "", "Jane", "jane@example.com", "", "Summary", ["Marketing"], ["Built dashboards"])
        payload = DocumentTailoringPayload(
            work_entries=[
                ResumeWorkEntry("Role", "2024", ["Managed optimized campaigns.", "Presented Presented that report.", "Created delivered both in-person."])
            ],
            key_skills=["Marketing"],
            cover_letter_text="Dear Hiring Team at Dow Jones,\n\nI am excited to apply for the role and bring experience in reporting, campaigns, and stakeholder coordination.\n\nSincerely,\nJane",
        )
        self.assertTrue(generator.should_escalate_tailoring(job, resume, payload))

    def test_validate_tailoring_payload_accepts_clean_payload(self) -> None:
        generator = DocumentGenerator()
        payload = DocumentTailoringPayload(
            work_entries=[
                ResumeWorkEntry("Role", "2024", ["Managed integrated campaigns and improved lead quality.", "Built dashboards for leadership reporting."])
            ],
            key_skills=["Marketing Strategy", "Dashboard Reporting"],
            cover_letter_text="Dear Hiring Team at Acme,\n\nI am excited to apply for the role and would bring experience in integrated campaigns, leadership reporting, and audience growth.\n\nSincerely,\nJane",
        )
        valid, issues = generator.validate_tailoring_payload(payload)
        self.assertTrue(valid)
        self.assertEqual(issues, [])

    def test_docx_generation_fails_when_template_sections_are_missing(self) -> None:
        with workspace_temp_dir() as tmp:
            from docx import Document

            bad_source = Path(tmp) / "bad_resume.docx"
            doc = Document()
            doc.add_paragraph("ALEX RIVERS")
            doc.add_paragraph("KEY SKILLS", style="Heading 1")
            doc.add_paragraph("Operations, Project Management")
            doc.save(str(bad_source))

            generator = DocumentGenerator()
            resume = parse_resume(bad_source, Path(tmp) / "bad_resume_cache.json")
            with self.assertRaises(ValueError):
                generator.generate(
                    Path(tmp),
                    Job("1", "Director", "Acme", "Remote", "", "Role", "board", "", "", "jobspy", "", ""),
                    resume,
                    MatchScore(0, "", [], [], False, "review", "", ""),
                    ai_notes="",
                )
          