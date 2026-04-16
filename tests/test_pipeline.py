from __future__ import annotations

import os
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from config import build_app_paths, load_or_create_config
from database import Database, Job
from document_tailoring import DocumentTailoringAttempt, DocumentTailoringPayload
from doc_generator import GeneratedDocs
from gmail_client import DeliveryResult
from match_scorer import StageEvaluation
from pipeline import JobBotPipeline
from portal_filler import PortalAutofillReadiness, PortalResult
from resume_parser import ResumeData, ResumeWorkEntry
from test_support import workspace_temp_dir


class PipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self._portal_readiness_patcher = patch(
            "pipeline.PortalFiller.check_readiness",
            return_value=PortalAutofillReadiness(
                ready=False,
                summary="Portal autofill: unavailable (test stub)",
                reason_code="test_stub",
                technical_detail="",
            ),
        )
        self._portal_readiness_patcher.start()
        self._doc_generate_patcher = patch("pipeline.DocumentGenerator.generate", new=self._fake_generate)
        self._doc_generate_patcher.start()

    def tearDown(self) -> None:
        self._doc_generate_patcher.stop()
        self._portal_readiness_patcher.stop()

    @staticmethod
    def _fake_generate(_generator, output_dir, job, resume, score, *, ai_notes="", tailoring_payload=None, progress_callback=None):
        if progress_callback:
            progress_callback("tailoring_resume", "Tailoring resume...", 3)
            progress_callback("exporting_pdf", "Exporting PDF...", 4)
            progress_callback("validating_pages", "Validating final page count...", 5)
            progress_callback("writing_cover_letter", "Writing cover letter...", 6)
        target_dir = Path(output_dir) / f"{job.employer}_{job.id[:8]}"
        target_dir.mkdir(parents=True, exist_ok=True)
        resume_pdf = target_dir / "resume.pdf"
        cover_txt = target_dir / "cover_letter.txt"
        cover_docx = target_dir / "cover_letter.docx"
        cover_pdf = target_dir / "cover_letter.pdf"
        resume_pdf.write_text("pdf", encoding="utf-8")
        cover_txt.write_text("cover", encoding="utf-8")
        cover_docx.write_text("docx", encoding="utf-8")
        cover_pdf.write_text("pdf", encoding="utf-8")
        return GeneratedDocs(
            output_dir=target_dir,
            resume_pdf_path=resume_pdf,
            cover_letter_pdf_path=cover_pdf,
            cover_letter_docx_path=cover_docx,
            cover_letter_txt_path=cover_txt,
            pdf_exporter_used="word_com",
            cover_letter_pdf_exporter_used="word_com",
            page_fit_attempts=1,
        )

    def test_prepare_document_tailoring_retries_only_cover_letter_when_resume_is_salvageable(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                local_payload = DocumentTailoringPayload(
                    work_entries=[ResumeWorkEntry("Role 1", "2024", ["Built dashboards for leadership reporting and campaign planning."])],
                    key_skills=["Dashboard Reporting"],
                    cover_letter_text="Dear Hiring Team at Acme,\n\nI am excited to apply for the Head of Marketing role and bring dashboard reporting and operational leadership experience.\n\nSincerely,\nJane",
                )
                first_ai = DocumentTailoringPayload(
                    work_entries=[ResumeWorkEntry("Role 1", "2024", ["Improved KPI visibility and reporting cadence for senior leadership teams."])],
                    key_skills=["Dashboard Reporting", "KPI Reporting"],
                    cover_letter_text="{'type': 'string'}",
                    route="openai",
                    ai_attempted=True,
                    provider="openai",
                    model="gpt-5-mini",
                )
                retry_ai = DocumentTailoringPayload(
                    work_entries=[ResumeWorkEntry("Role 1", "2024", ["Changed but should be ignored."])],
                    key_skills=["Changed"],
                    cover_letter_text="Dear Hiring Team at Acme,\n\nI am excited to apply for the Head of Marketing role and would bring experience in KPI visibility, reporting, and cross-functional execution.\n\nSincerely,\nJane",
                    route="openai",
                    ai_attempted=True,
                    provider="openai",
                    model="gpt-5-mini",
                )
                resume = ResumeData(
                    "", "", "Jane", "jane@example.com", "", "Summary", ["Dashboard Reporting"], ["Built dashboards"],
                    work_experience_entries=local_payload.work_entries,
                )
                job = Job("1", "Head of Marketing", "Acme", "Remote", "", "Lead B2B growth", "board", "", "", "jobspy", "", "")
                score = StageEvaluation(
                    "strong", "openai", "gpt-5-mini", "scored", "review", 80, 0.8, "Good fit", [], [], 0, 0, 0.0, "", ""
                ).to_match_score(config.final_apply_threshold)
                pipeline.doc_generator.build_local_tailoring_payload = lambda *args, **kwargs: local_payload
                pipeline.doc_generator.should_escalate_tailoring = lambda *args, **kwargs: True
                requested_sections: list[set[str] | None] = []

                def fake_tailor(*args, **kwargs):
                    requested_sections.append(kwargs.get("requested_sections"))
                    if kwargs.get("retry_count", 0) == 0:
                        return DocumentTailoringAttempt(attempted=True, provider="openai", model="gpt-5-mini", retry_count=0, payload=first_ai)
                    return DocumentTailoringAttempt(attempted=True, provider="openai", model="gpt-5-mini", retry_count=1, payload=retry_ai)

                pipeline.scorer.tailor_documents_with_ai = fake_tailor
                payload = pipeline._prepare_document_tailoring(job, resume, score, ai_notes="")
                self.assertEqual(requested_sections[1], {"cover_letter"})
                self.assertEqual(payload.route, "openai")
                self.assertEqual(payload.resume_ai_status, "accepted")
                self.assertEqual(payload.cover_letter_ai_status, "accepted")
                self.assertFalse(payload.resume_retry_performed)
                self.assertTrue(payload.cover_letter_retry_performed)
                self.assertEqual(payload.work_entries[0].bullets[0], "Improved KPI visibility and reporting cadence for senior leadership teams.")

    def test_prepare_document_tailoring_uses_best_partial_when_retry_fails(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                local_payload = DocumentTailoringPayload(
                    work_entries=[ResumeWorkEntry("Role 1", "2024", ["Managed and optimized campaigns.", "Presented the report to leadership."])],
                    key_skills=["Dashboard Reporting"],
                    cover_letter_text="Dear Hiring Team at Acme,\n\nI am excited to apply for the role and bring dashboard reporting and operational leadership experience.\n\nSincerely,\nJane",
                )
                first_ai = DocumentTailoringPayload(
                    work_entries=[ResumeWorkEntry("Role 1", "2024", ["Improved KPI visibility and reporting cadence.", "Presented Presented that report."])],
                    key_skills=["Dashboard Reporting"],
                    cover_letter_text="Dear Hiring Team at Acme,\n\nI am excited to apply for the role and would bring experience in reporting and cross-functional execution.\n\nSincerely,\nJane",
                    route="openai",
                    ai_attempted=True,
                    provider="openai",
                    model="gpt-5-mini",
                )
                resume = ResumeData(
                    "", "", "Jane", "jane@example.com", "", "Summary", ["Dashboard Reporting"], ["Built dashboards"],
                    work_experience_entries=local_payload.work_entries,
                )
                job = Job("1", "Head of Marketing", "Acme", "Remote", "", "Lead B2B growth", "board", "", "", "jobspy", "", "")
                score = StageEvaluation(
                    "strong", "openai", "gpt-5-mini", "scored", "review", 80, 0.8, "Good fit", [], [], 0, 0, 0.0, "", ""
                ).to_match_score(config.final_apply_threshold)
                pipeline.doc_generator.build_local_tailoring_payload = lambda *args, **kwargs: local_payload
                pipeline.doc_generator.should_escalate_tailoring = lambda *args, **kwargs: True

                def fake_tailor(*args, **kwargs):
                    if kwargs.get("retry_count", 0) == 0:
                        return DocumentTailoringAttempt(attempted=True, provider="openai", model="gpt-5-mini", retry_count=0, payload=first_ai)
                    return DocumentTailoringAttempt(attempted=True, provider="openai", model="gpt-5-mini", failure_reason="OpenAI request failed: timeout", retry_count=1)

                pipeline.scorer.tailor_documents_with_ai = fake_tailor
                payload = pipeline._prepare_document_tailoring(job, resume, score, ai_notes="")
                self.assertEqual(payload.route, "openai")
                self.assertEqual(payload.resume_ai_status, "partial")
                self.assertEqual(payload.cover_letter_ai_status, "accepted")
                self.assertEqual(payload.rejected_bullets_repaired, 1)
                self.assertEqual(payload.work_entries[0].bullets[0], "Improved KPI visibility and reporting cadence.")
                self.assertEqual(payload.work_entries[0].bullets[1], "Presented the report to leadership.")

    def test_prepare_document_tailoring_emits_section_retry_progress_messages(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                local_payload = DocumentTailoringPayload(
                    work_entries=[ResumeWorkEntry("Role 1", "2024", ["Built dashboards for leadership reporting and campaign planning."])],
                    key_skills=["Dashboard Reporting"],
                    cover_letter_text="Dear Hiring Team at Acme,\n\nI am excited to apply for the role and bring dashboard reporting experience.\n\nSincerely,\nJane",
                )
                first_ai = DocumentTailoringPayload(
                    work_entries=[ResumeWorkEntry("Role 1", "2024", ["Built dashboards for leadership reporting and campaign planning."])],
                    key_skills=["Dashboard Reporting"],
                    cover_letter_text="{'type': 'string'}",
                    route="openai",
                    ai_attempted=True,
                    provider="openai",
                    model="gpt-5-mini",
                )
                retry_ai = DocumentTailoringPayload(
                    work_entries=[ResumeWorkEntry("Role 1", "2024", ["Ignore this retry bullet because resume changes should be preserved."])],
                    key_skills=["Dashboard Reporting"],
                    cover_letter_text="Dear Hiring Team at Acme,\n\nI am excited to apply for the role and bring dashboard reporting experience.\n\nSincerely,\nJane",
                    route="openai",
                    ai_attempted=True,
                    provider="openai",
                    model="gpt-5-mini",
                )
                resume = ResumeData(
                    "", "", "Jane", "jane@example.com", "", "Summary", ["Dashboard Reporting"], ["Built dashboards"],
                    work_experience_entries=local_payload.work_entries,
                )
                job = Job("1", "Head of Marketing", "Acme", "Remote", "", "Lead B2B growth", "board", "", "", "jobspy", "", "")
                score = StageEvaluation(
                    "strong", "openai", "gpt-5-mini", "scored", "review", 80, 0.8, "Good fit", [], [], 0, 0, 0.0, "", ""
                ).to_match_score(config.final_apply_threshold)
                pipeline.doc_generator.build_local_tailoring_payload = lambda *args, **kwargs: local_payload
                pipeline.doc_generator.should_escalate_tailoring = lambda *args, **kwargs: True
                pipeline.scorer.tailor_documents_with_ai = lambda *args, **kwargs: (
                    DocumentTailoringAttempt(attempted=True, provider="openai", model="gpt-5-mini", retry_count=0, payload=first_ai)
                    if kwargs.get("retry_count", 0) == 0
                    else DocumentTailoringAttempt(attempted=True, provider="openai", model="gpt-5-mini", retry_count=1, payload=retry_ai)
                )
                events: list[tuple[str, str, int]] = []
                pipeline._prepare_document_tailoring(job, resume, score, ai_notes="", progress_callback=lambda stage, message, progress: events.append((stage, message, progress)))
                messages = [message for _stage, message, _progress in events]
                self.assertIn("Validating AI content...", messages)
                self.assertIn("Repairing invalid AI content...", messages)
                self.assertIn("Retrying cover letter generation...", messages)

    def test_prepare_document_tailoring_records_specific_parse_failure_and_retry_count(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                local_payload = DocumentTailoringPayload(
                    work_entries=[ResumeWorkEntry("Role 1", "2024", ["Built dashboards"])],
                    key_skills=["Dashboard Reporting"],
                    cover_letter_text="Local",
                )
                resume = ResumeData(
                    "", "", "Jane", "jane@example.com", "", "Summary", ["Dashboard Reporting"], ["Built dashboards"],
                    work_experience_entries=local_payload.work_entries,
                )
                job = Job("1", "Head of Marketing", "Acme", "Remote", "", "Lead B2B growth", "board", "", "", "jobspy", "", "")
                score = StageEvaluation(
                    "strong", "openai", "gpt-5-mini", "scored", "review", 80, 0.8, "Good fit", [], [], 0, 0, 0.0, "", ""
                ).to_match_score(config.final_apply_threshold)
                pipeline.doc_generator.build_local_tailoring_payload = lambda *args, **kwargs: local_payload
                pipeline.doc_generator.should_escalate_tailoring = lambda *args, **kwargs: True
                pipeline.scorer.tailor_documents_with_ai = lambda *args, **kwargs: DocumentTailoringAttempt(
                    attempted=True,
                    provider="openai",
                    model="gpt-5-mini",
                    failure_reason="OpenAI response could not be parsed as JSON.",
                    retry_count=kwargs.get("retry_count", 0),
                )

                payload = pipeline._prepare_document_tailoring(job, resume, score, ai_notes="")
                self.assertEqual(payload.route, "fallback")
                self.assertTrue(payload.ai_attempted)
                self.assertEqual(payload.provider, "openai")
                self.assertEqual(payload.model, "gpt-5-mini")
                self.assertEqual(payload.retry_count, 1)
                self.assertIn("parsed as json", payload.fallback_reason.lower())

    def test_pipeline_records_completed_run(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, SQL\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                config.cheap_stage_provider = "openai"
                config.strong_stage_provider = "openai"
                config.openai_api_key = "test-key"
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                sample_job = Job(
                    id="1",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Build APIs",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="usajobs",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(sample_job, {"sample": True})]
                pipeline.scorer.cheap_evaluate = lambda job, resume, force_escalate=False: StageEvaluation(
                    "cheap", "openai", "gpt-5-nano", "scored", "escalate", 72, 0.8, "Promising", ["Python"], ["AWS"], 100, 10, 0.0001, "", "2026-01-01T00:00:00+00:00"
                )
                pipeline.scorer.strong_evaluate = lambda job, resume: StageEvaluation(
                    "strong", "openai", "gpt-5-mini", "scored", "apply_candidate", 91, 0.9, "Great fit", ["Python"], ["AWS"], 200, 20, 0.0002, "", "2026-01-01T00:00:00+00:00"
                )
                pipeline.scorer.document_generation_notes = lambda job, resume, strong_eval: "Tailored note"
                pipeline.gmail_client.deliver_match = lambda job, score, docs, **kwargs: type(
                    "Delivery",
                    (),
                    {"method": "local", "status": "skipped", "message_id": "", "error_message": ""},
                )()
                result = pipeline.run()
                self.assertEqual(result.status, "completed")
                rows = database.list_review_rows()
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0]["delivery_status"], "pending_approval")
                self.assertEqual(rows[0]["document_status"], "pending")

    def test_auto_mode_still_sends_when_enabled(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "auto"
                config.cheap_stage_provider = "openai"
                config.strong_stage_provider = "openai"
                config.openai_api_key = "test-key"
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, SQL\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                sample_job = Job(
                    id="2",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Email your resume to hiring@example.com",
                    apply_method="email",
                    apply_url="https://example.com",
                    hiring_manager_email="hiring@example.com",
                    source="usajobs",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(sample_job, {"sample": True})]
                pipeline.scorer.cheap_evaluate = lambda job, resume, force_escalate=False: StageEvaluation(
                    "cheap", "openai", "gpt-5-nano", "scored", "escalate", 72, 0.8, "Promising", ["Python"], ["AWS"], 100, 10, 0.0001, "", "2026-01-01T00:00:00+00:00"
                )
                pipeline.scorer.strong_evaluate = lambda job, resume: StageEvaluation(
                    "strong", "openai", "gpt-5-mini", "scored", "apply_candidate", 91, 0.9, "Great fit", ["Python"], ["AWS"], 200, 20, 0.0002, "", "2026-01-01T00:00:00+00:00"
                )
                pipeline.scorer.document_generation_notes = lambda job, resume, strong_eval: "Tailored note"
                pipeline.gmail_client.deliver_match = lambda job, score, docs, **kwargs: type(
                    "Delivery",
                    (),
                    {"method": "gmail_employer", "status": "sent", "message_id": "abc123", "error_message": ""},
                )()
                pipeline.run()
                row = database.get_review_row("2")
                self.assertEqual(row["delivery_status"], "sent")
                self.assertEqual(row["delivery_method"], "gmail_employer")
                self.assertEqual(row["apply_method"], "email")
                self.assertEqual(row["hiring_manager_email"], "hiring@example.com")

    def test_auto_mode_keeps_board_jobs_in_review_queue(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "auto"
                config.cheap_stage_provider = "openai"
                config.strong_stage_provider = "openai"
                config.openai_api_key = "test-key"
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, SQL\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                sample_job = Job(
                    id="3",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Apply on the company site.",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="usajobs",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(sample_job, {"sample": True})]
                pipeline.scorer.cheap_evaluate = lambda job, resume, force_escalate=False: StageEvaluation(
                    "cheap", "openai", "gpt-5-nano", "scored", "escalate", 72, 0.8, "Promising", ["Python"], ["AWS"], 100, 10, 0.0001, "", "2026-01-01T00:00:00+00:00"
                )
                pipeline.scorer.strong_evaluate = lambda job, resume: StageEvaluation(
                    "strong", "openai", "gpt-5-mini", "scored", "apply_candidate", 91, 0.9, "Great fit", ["Python"], ["AWS"], 200, 20, 0.0002, "", "2026-01-01T00:00:00+00:00"
                )
                pipeline.scorer.document_generation_notes = lambda job, resume, strong_eval: "Tailored note"
                pipeline.gmail_client.deliver_match = lambda job, score, docs, **kwargs: type(
                    "Delivery",
                    (),
                    {"method": "gmail", "status": "sent", "message_id": "abc123", "error_message": ""},
                )()
                pipeline.run()
                row = database.get_review_row("3")
                self.assertEqual(row["delivery_status"], "pending_approval")
                self.assertEqual(row["apply_method"], "board")

    def test_approve_and_send_employer_email_uses_gmail_without_fallback(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Marketing leader\nSkills: Marketing\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                database.upsert_job(
                    Job(
                        id="approve-email-1",
                        title="Head of Marketing",
                        employer="Acme",
                        location="Remote",
                        salary_range="",
                        description_full="Email your resume to hiring@example.com",
                        apply_method="email",
                        apply_url="https://example.com",
                        hiring_manager_email="hiring@example.com",
                        source="jobspy",
                        posted_at="2026-01-01T00:00:00+00:00",
                        scraped_at="2026-01-01T00:00:00+00:00",
                    ),
                    {},
                )
                database.record_match_result(
                    "approve-email-1",
                    score=82,
                    rationale="Strong fit",
                    strengths=[],
                    gaps=[],
                    is_match=True,
                    status="review",
                    error_message="",
                    scored_at="2026-01-01T00:00:00+00:00",
                )
                out = Path(tmp) / "output"
                out.mkdir(parents=True, exist_ok=True)
                resume_pdf = out / "resume.pdf"
                cover = out / "cover_letter.pdf"
                resume_pdf.write_text("pdf", encoding="utf-8")
                cover.write_text("pdf", encoding="utf-8")
                database.record_generated_documents(
                    "approve-email-1",
                    output_dir=str(out),
                    resume_docx_path="",
                    resume_pdf_path=str(resume_pdf),
                    cover_letter_pdf_path=str(cover),
                    cover_letter_path="",
                    status="generated",
                    error_message="",
                    generated_at="2026-01-01T00:00:00+00:00",
                )
                captured: dict[str, object] = {}

                def fake_deliver(job, score, docs, **kwargs):
                    captured["allow_fallback"] = kwargs.get("allow_fallback")
                    return DeliveryResult("gmail_employer", "sent", "abc123", "")

                pipeline.gmail_client.deliver_match = fake_deliver
                result = pipeline.approve_and_send("approve-email-1")
                self.assertEqual(result.status, "sent")
                self.assertEqual(result.method, "gmail_employer")
                self.assertFalse(captured["allow_fallback"])
                row = database.get_review_row("approve-email-1")
                self.assertEqual(row["delivery_status"], "sent")
                self.assertEqual(row["delivery_method"], "gmail_employer")

    def test_approve_and_send_email_failure_records_approval_log(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                config.gmail.enabled = True
                config.gmail.sender_email = "robert@example.com"
                config.gmail.client_secrets_file = str(Path(tmp) / "client_secret.json")
                Path(config.gmail.client_secrets_file).write_text("{}", encoding="utf-8")
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Marketing leader\nSkills: Marketing\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                database.upsert_job(
                    Job(
                        id="approve-email-fail-1",
                        title="VP Communications",
                        employer="Acme",
                        location="Remote",
                        salary_range="",
                        description_full="Email your resume to recruiting@example.com",
                        apply_method="email",
                        apply_url="https://example.com",
                        hiring_manager_email="recruiting@example.com",
                        source="jobspy",
                        posted_at="2026-01-01T00:00:00+00:00",
                        scraped_at="2026-01-01T00:00:00+00:00",
                    ),
                    {},
                )
                database.record_match_result(
                    "approve-email-fail-1",
                    score=82,
                    rationale="Strong fit",
                    strengths=[],
                    gaps=[],
                    is_match=True,
                    status="review",
                    error_message="",
                    scored_at="2026-01-01T00:00:00+00:00",
                )
                out = Path(tmp) / "output"
                out.mkdir(parents=True, exist_ok=True)
                resume_pdf = out / "resume.pdf"
                cover = out / "cover_letter.pdf"
                resume_pdf.write_text("pdf", encoding="utf-8")
                cover.write_text("pdf", encoding="utf-8")
                database.record_generated_documents(
                    "approve-email-fail-1",
                    output_dir=str(out),
                    resume_docx_path="",
                    resume_pdf_path=str(resume_pdf),
                    cover_letter_pdf_path=str(cover),
                    cover_letter_path="",
                    status="generated",
                    error_message="",
                    generated_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.gmail_client.deliver_match = lambda *args, **kwargs: DeliveryResult(
                    "gmail_employer",
                    "failed",
                    "",
                    "Google OAuth blocked access; app is still in testing and this account must be added as a test user.",
                )
                result = pipeline.approve_and_send("approve-email-fail-1")
                self.assertEqual(result.status, "failed")
                self.assertEqual(result.method, "gmail_employer")
                row = database.get_review_row("approve-email-fail-1")
                self.assertEqual(row["delivery_status"], "failed")
                self.assertEqual(row["delivery_method"], "gmail_employer")
                self.assertIn("Route: email", row["approval_log"])
                self.assertIn("Google OAuth blocked access", row["approval_log"])

    def test_approve_and_send_email_without_employer_email_records_failure(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Marketing leader\nSkills: Marketing\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                database.upsert_job(
                    Job(
                        id="approve-email-fail-2",
                        title="VP Communications",
                        employer="Acme",
                        location="Remote",
                        salary_range="",
                        description_full="Apply by email.",
                        apply_method="email",
                        apply_url="https://example.com",
                        hiring_manager_email="",
                        source="jobspy",
                        posted_at="2026-01-01T00:00:00+00:00",
                        scraped_at="2026-01-01T00:00:00+00:00",
                    ),
                    {},
                )
                database.record_match_result(
                    "approve-email-fail-2",
                    score=82,
                    rationale="Strong fit",
                    strengths=[],
                    gaps=[],
                    is_match=True,
                    status="review",
                    error_message="",
                    scored_at="2026-01-01T00:00:00+00:00",
                )
                out = Path(tmp) / "output"
                out.mkdir(parents=True, exist_ok=True)
                resume_pdf = out / "resume.pdf"
                cover = out / "cover_letter.pdf"
                resume_pdf.write_text("pdf", encoding="utf-8")
                cover.write_text("pdf", encoding="utf-8")
                database.record_generated_documents(
                    "approve-email-fail-2",
                    output_dir=str(out),
                    resume_docx_path="",
                    resume_pdf_path=str(resume_pdf),
                    cover_letter_pdf_path=str(cover),
                    cover_letter_path="",
                    status="generated",
                    error_message="",
                    generated_at="2026-01-01T00:00:00+00:00",
                )
                result = pipeline.approve_and_send("approve-email-fail-2")
                self.assertEqual(result.status, "blocked_hr_email")
                self.assertEqual(result.method, "manual")
                self.assertIn("No validated HR/application email was found", result.error_message)
                row = database.get_review_row("approve-email-fail-2")
                self.assertIn("HR email: not present", row["approval_log"])
                self.assertIn("No validated HR/application email was found", row["approval_log"])

    def test_approve_and_send_blocks_legacy_email_row_with_compliance_recipient(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                config.gmail.enabled = True
                config.gmail.sender_email = "robert@example.com"
                config.gmail.client_secrets_file = str(Path(tmp) / "client_secret.json")
                Path(config.gmail.client_secrets_file).write_text("{}", encoding="utf-8")
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Marketing leader\nSkills: Marketing\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                database.upsert_job(
                    Job(
                        id="approve-email-block-1",
                        title="VP Communications",
                        employer="Acme",
                        location="Remote",
                        salary_range="",
                        description_full=(
                            "For questions about privacy, contact compliance@acme.com.\n"
                            "This inbox is not monitored for application status updates."
                        ),
                        apply_method="email",
                        apply_url="https://example.com",
                        hiring_manager_email="compliance@acme.com",
                        source="jobspy",
                        posted_at="2026-01-01T00:00:00+00:00",
                        scraped_at="2026-01-01T00:00:00+00:00",
                    ),
                    {},
                )
                database.record_match_result(
                    "approve-email-block-1",
                    score=82,
                    rationale="Strong fit",
                    strengths=[],
                    gaps=[],
                    is_match=True,
                    status="review",
                    error_message="",
                    scored_at="2026-01-01T00:00:00+00:00",
                )
                out = Path(tmp) / "output"
                out.mkdir(parents=True, exist_ok=True)
                resume_pdf = out / "resume.pdf"
                cover = out / "cover_letter.pdf"
                resume_pdf.write_text("pdf", encoding="utf-8")
                cover.write_text("pdf", encoding="utf-8")
                database.record_generated_documents(
                    "approve-email-block-1",
                    output_dir=str(out),
                    resume_docx_path="",
                    resume_pdf_path=str(resume_pdf),
                    cover_letter_pdf_path=str(cover),
                    cover_letter_path="",
                    status="generated",
                    error_message="",
                    generated_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.gmail_client.deliver_match = lambda *args, **kwargs: (_ for _ in ()).throw(
                    AssertionError("gmail should not send blocked recipients")
                )
                result = pipeline.approve_and_send("approve-email-block-1")
                self.assertEqual(result.status, "blocked_hr_email")
                self.assertEqual(result.method, "manual")
                self.assertIn("validated hr/application email", result.error_message.lower())
                row = database.get_review_row("approve-email-block-1")
                self.assertEqual(row["delivery_status"], "blocked_hr_email")
                self.assertEqual(row["delivery_method"], "manual")
                self.assertIn("HR email: not present", row["approval_log"])
                self.assertIn("Validated HR recipient: Not present", row["approval_log"])

    def test_approve_and_send_supported_portal_submits_without_gmail(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Marketing leader\nSkills: Marketing\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                database.upsert_job(
                    Job(
                        id="approve-portal-1",
                        title="Head of Marketing",
                        employer="Acme",
                        location="Remote",
                        salary_range="",
                        description_full="Apply on company site.",
                        apply_method="board",
                        apply_url="https://boards.greenhouse.io/acme/jobs/123",
                        hiring_manager_email="",
                        source="jobspy",
                        posted_at="2026-01-01T00:00:00+00:00",
                        scraped_at="2026-01-01T00:00:00+00:00",
                    ),
                    {},
                )
                database.record_match_result(
                    "approve-portal-1",
                    score=82,
                    rationale="Strong fit",
                    strengths=[],
                    gaps=[],
                    is_match=True,
                    status="review",
                    error_message="",
                    scored_at="2026-01-01T00:00:00+00:00",
                )
                out = Path(tmp) / "output"
                out.mkdir(parents=True, exist_ok=True)
                resume_pdf = out / "resume.pdf"
                cover = out / "cover_letter.pdf"
                resume_pdf.write_text("pdf", encoding="utf-8")
                cover.write_text("pdf", encoding="utf-8")
                database.record_generated_documents(
                    "approve-portal-1",
                    output_dir=str(out),
                    resume_docx_path="",
                    resume_pdf_path=str(resume_pdf),
                    cover_letter_pdf_path=str(cover),
                    cover_letter_path="",
                    status="generated",
                    error_message="",
                    generated_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.gmail_client.deliver_match = lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("gmail should not be used for board jobs"))
                pipeline.portal_readiness.ready = True
                pipeline.portal_readiness.summary = "Portal autofill: ready"
                pipeline.portal_readiness.reason_code = "ready"
                pipeline.portal_filler.fill = lambda *args, **kwargs: PortalResult("submitted", "Submitted via Greenhouse.", "greenhouse")
                result = pipeline.approve_and_send("approve-portal-1")
                self.assertEqual(result.status, "submitted")
                self.assertEqual(result.method, "greenhouse")
                row = database.get_review_row("approve-portal-1")
                self.assertEqual(row["delivery_status"], "submitted")
                self.assertEqual(row["delivery_method"], "greenhouse")

    def test_approve_and_send_unsupported_portal_opens_manual_page(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Marketing leader\nSkills: Marketing\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                database.upsert_job(
                    Job(
                        id="approve-portal-2",
                        title="Head of Marketing",
                        employer="Acme",
                        location="Remote",
                        salary_range="",
                        description_full="Apply on company site.",
                        apply_method="board",
                        apply_url="https://careers.example.com/jobs/123",
                        hiring_manager_email="",
                        source="jobspy",
                        posted_at="2026-01-01T00:00:00+00:00",
                        scraped_at="2026-01-01T00:00:00+00:00",
                    ),
                    {},
                )
                database.record_match_result(
                    "approve-portal-2",
                    score=82,
                    rationale="Strong fit",
                    strengths=[],
                    gaps=[],
                    is_match=True,
                    status="review",
                    error_message="",
                    scored_at="2026-01-01T00:00:00+00:00",
                )
                out = Path(tmp) / "output"
                out.mkdir(parents=True, exist_ok=True)
                resume_pdf = out / "resume.pdf"
                cover = out / "cover_letter.pdf"
                resume_pdf.write_text("pdf", encoding="utf-8")
                cover.write_text("pdf", encoding="utf-8")
                database.record_generated_documents(
                    "approve-portal-2",
                    output_dir=str(out),
                    resume_docx_path="",
                    resume_pdf_path=str(resume_pdf),
                    cover_letter_pdf_path=str(cover),
                    cover_letter_path="",
                    status="generated",
                    error_message="",
                    generated_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.portal_readiness.ready = True
                pipeline.portal_readiness.summary = "Portal autofill: ready"
                pipeline.portal_readiness.reason_code = "ready"
                pipeline.portal_filler.fill = lambda *args, **kwargs: PortalResult("unsupported", "Unsupported portal.", "unknown")
                with patch("pipeline.webbrowser.open", return_value=True) as open_mock:
                    result = pipeline.approve_and_send("approve-portal-2")
                self.assertEqual(result.status, "unsupported_portal")
                self.assertEqual(result.method, "portal")
                self.assertIn("opened apply page", result.error_message.lower())
                open_mock.assert_called_once()
                row = database.get_review_row("approve-portal-2")
                self.assertEqual(row["delivery_status"], "unsupported_portal")
                self.assertEqual(row["delivery_method"], "portal")

    def test_approve_and_send_indeed_posting_page_preserves_indeed_platform_in_log(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Marketing leader\nSkills: Marketing\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                database.upsert_job(
                    Job(
                        id="approve-indeed-1",
                        title="Head of Marketing",
                        employer="Acme",
                        location="Remote",
                        salary_range="",
                        description_full="Apply on Indeed.",
                        apply_method="board",
                        apply_url="https://www.indeed.com/viewjob?jk=123",
                        hiring_manager_email="",
                        source="jobspy",
                        posted_at="2026-01-01T00:00:00+00:00",
                        scraped_at="2026-01-01T00:00:00+00:00",
                    ),
                    {},
                )
                database.record_match_result(
                    "approve-indeed-1",
                    score=82,
                    rationale="Strong fit",
                    strengths=[],
                    gaps=[],
                    is_match=True,
                    status="review",
                    error_message="",
                    scored_at="2026-01-01T00:00:00+00:00",
                )
                out = Path(tmp) / "output"
                out.mkdir(parents=True, exist_ok=True)
                resume_pdf = out / "resume.pdf"
                cover = out / "cover_letter.pdf"
                resume_pdf.write_text("pdf", encoding="utf-8")
                cover.write_text("pdf", encoding="utf-8")
                database.record_generated_documents(
                    "approve-indeed-1",
                    output_dir=str(out),
                    resume_docx_path="",
                    resume_pdf_path=str(resume_pdf),
                    cover_letter_pdf_path=str(cover),
                    cover_letter_path="",
                    status="generated",
                    error_message="",
                    generated_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.portal_readiness.ready = True
                pipeline.portal_readiness.summary = "Portal autofill: ready"
                pipeline.portal_readiness.reason_code = "ready"
                pipeline.portal_filler.fill = lambda *args, **kwargs: PortalResult(
                    "unsupported",
                    "Indeed posting page is not an automatable Indeed apply form.",
                    "indeed",
                )
                with patch("pipeline.webbrowser.open", return_value=True):
                    result = pipeline.approve_and_send("approve-indeed-1")
                self.assertEqual(result.status, "unsupported_portal")
                self.assertEqual(result.method, "indeed")
                self.assertIn("indeed posting page", result.error_message.lower())
                row = database.get_review_row("approve-indeed-1")
                self.assertEqual(row["delivery_method"], "indeed")
                self.assertIn("Portal platform: indeed", row["approval_log"])
                self.assertIn("Fallback: indeed posting page", row["approval_log"])

    def test_approve_and_send_screening_questions_open_manual_review(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Marketing leader\nSkills: Marketing\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                database.upsert_job(
                    Job(
                        id="approve-indeed-2",
                        title="Head of Marketing",
                        employer="Acme",
                        location="Remote",
                        salary_range="",
                        description_full="Apply on Indeed.",
                        apply_method="board",
                        apply_url="https://apply.indeed.com/indeedapply/form/abc",
                        hiring_manager_email="",
                        source="jobspy",
                        posted_at="2026-01-01T00:00:00+00:00",
                        scraped_at="2026-01-01T00:00:00+00:00",
                    ),
                    {},
                )
                database.record_match_result(
                    "approve-indeed-2",
                    score=82,
                    rationale="Strong fit",
                    strengths=[],
                    gaps=[],
                    is_match=True,
                    status="review",
                    error_message="",
                    scored_at="2026-01-01T00:00:00+00:00",
                )
                out = Path(tmp) / "output"
                out.mkdir(parents=True, exist_ok=True)
                resume_pdf = out / "resume.pdf"
                cover = out / "cover_letter.pdf"
                resume_pdf.write_text("pdf", encoding="utf-8")
                cover.write_text("pdf", encoding="utf-8")
                database.record_generated_documents(
                    "approve-indeed-2",
                    output_dir=str(out),
                    resume_docx_path="",
                    resume_pdf_path=str(resume_pdf),
                    cover_letter_pdf_path=str(cover),
                    cover_letter_path="",
                    status="generated",
                    error_message="",
                    generated_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.portal_readiness.ready = True
                pipeline.portal_readiness.summary = "Portal autofill: ready"
                pipeline.portal_readiness.reason_code = "ready"
                pipeline.portal_filler.fill = lambda *args, **kwargs: PortalResult(
                    "screening_questions",
                    "Indeed apply flow includes screening questions that require manual review.",
                    "indeed",
                )
                with patch("pipeline.webbrowser.open", return_value=True):
                    result = pipeline.approve_and_send("approve-indeed-2")
                self.assertEqual(result.status, "opened_manual")
                self.assertEqual(result.method, "indeed")
                self.assertIn("screening questions", result.error_message.lower())
                row = database.get_review_row("approve-indeed-2")
                self.assertIn("Fallback: screening questions", row["approval_log"])

    def test_approve_and_send_skips_portal_when_readiness_unavailable(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Marketing leader\nSkills: Marketing\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                pipeline.portal_readiness.ready = False
                pipeline.portal_readiness.summary = "Portal autofill: unavailable (Playwright not installed)"
                pipeline.portal_readiness.reason_code = "missing_playwright"
                pipeline.portal_readiness.technical_detail = "ImportError: No module named 'playwright'"
                database.upsert_job(
                    Job(
                        id="approve-portal-3",
                        title="Head of Marketing",
                        employer="Acme",
                        location="Remote",
                        salary_range="",
                        description_full="Apply on company site.",
                        apply_method="board",
                        apply_url="https://careers.example.com/jobs/456",
                        hiring_manager_email="",
                        source="jobspy",
                        posted_at="2026-01-01T00:00:00+00:00",
                        scraped_at="2026-01-01T00:00:00+00:00",
                    ),
                    {},
                )
                database.record_match_result(
                    "approve-portal-3",
                    score=82,
                    rationale="Strong fit",
                    strengths=[],
                    gaps=[],
                    is_match=True,
                    status="review",
                    error_message="",
                    scored_at="2026-01-01T00:00:00+00:00",
                )
                out = Path(tmp) / "output"
                out.mkdir(parents=True, exist_ok=True)
                resume_pdf = out / "resume.pdf"
                cover = out / "cover_letter.pdf"
                resume_pdf.write_text("pdf", encoding="utf-8")
                cover.write_text("pdf", encoding="utf-8")
                database.record_generated_documents(
                    "approve-portal-3",
                    output_dir=str(out),
                    resume_docx_path="",
                    resume_pdf_path=str(resume_pdf),
                    cover_letter_pdf_path=str(cover),
                    cover_letter_path="",
                    status="generated",
                    error_message="",
                    generated_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.portal_filler.fill = lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("portal filler should not run"))
                with patch("pipeline.webbrowser.open", return_value=True) as open_mock:
                    result = pipeline.approve_and_send("approve-portal-3")
                self.assertEqual(result.status, "opened_manual")
                self.assertEqual(result.method, "portal")
                self.assertIn("portal autofill unavailable", result.error_message.lower())
                open_mock.assert_called_once()
                row = database.get_review_row("approve-portal-3")
                self.assertIn("Portal readiness detail: ImportError", row["approval_log"])

    def test_verify_portal_autofill_runtime_returns_readiness_and_emits_progress(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                pipeline.portal_filler.check_readiness = lambda: type(
                    "Readiness",
                    (),
                    {
                        "ready": False,
                        "summary": "Portal autofill: unavailable (Chromium launch failed)",
                        "reason_code": "browser_launch_failed",
                        "technical_detail": "winerror_5_access_denied: PermissionError: [WinError 5] Access is denied",
                    },
                )()
                events: list[tuple[str, str, int]] = []
                readiness = pipeline.verify_portal_autofill_runtime(
                    progress_callback=lambda stage, message, progress: events.append((stage, message, progress))
                )
                self.assertFalse(readiness.ready)
                self.assertEqual(readiness.reason_code, "browser_launch_failed")
                self.assertEqual(events[0], ("starting", "Testing portal autofill runtime...", 0))
                self.assertEqual(events[1], ("checking_runtime", "Importing Playwright and launching headless Chromium...", 4))
                self.assertEqual(events[2][0], "runtime_checked")
                self.assertIn("winerror_5_access_denied", events[2][1])

    def test_semi_auto_can_skip_ai_scoring_and_queue_fast(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, SQL\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                sample_job = Job(
                    id="4",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Build APIs",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="jobspy",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(sample_job, {"sample": True})]
                pipeline.scorer.cheap_evaluate = lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("cheap scorer should not run"))
                pipeline.scorer.strong_evaluate = lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("strong scorer should not run"))
                result = pipeline.run()
                self.assertEqual(result.status, "completed")
                self.assertEqual(result.jobs_matched, 1)
                row = database.get_review_row("4")
                self.assertEqual(row["match_status"], "review")
                self.assertEqual(row["delivery_status"], "pending_approval")
                self.assertEqual(row["document_status"], "pending")

    def test_pipeline_records_structured_run_metrics(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.progressive_queue_enabled = True
                config.cheap_ai_top_n = 1
                config.strong_ai_top_n = 1
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, APIs\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                sample_job = Job(
                    id="metrics-1",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Build APIs with Python",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="jobspy",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(sample_job, {"sample": True})]
                pipeline.scorer.cheap_evaluate = lambda *args, **kwargs: StageEvaluation(
                    "cheap", "openai", "gpt-5-nano", "scored", "review", 78, 0.8, "Promising fit", ["Python"], [], 10, 5, 0.001, "", datetime.now(timezone.utc).isoformat()
                )
                pipeline.scorer.strong_evaluate = lambda *args, **kwargs: StageEvaluation(
                    "strong", "openai", "gpt-5-mini", "scored", "review", 82, 0.85, "Worth review", ["Python"], [], 12, 6, 0.01, "", datetime.now(timezone.utc).isoformat()
                )
                result = pipeline.run()
                self.assertEqual(result.status, "completed")
                latest_run = database.latest_run()
                self.assertIsNotNone(latest_run)
                assert latest_run is not None
                self.assertEqual(latest_run["jobs_seen"], 1)
                self.assertEqual(latest_run["fast_rank_candidates"], 1)
                self.assertEqual(latest_run["fast_rank_survivors"], 1)
                self.assertEqual(latest_run["cheap_shortlist_size"], 1)
                self.assertEqual(latest_run["cheap_stage_provider"], "openai")
                self.assertEqual(latest_run["strong_shortlist_size"], 1)
                self.assertEqual(latest_run["strong_stage_model"], "gpt-5-mini")

    def test_pipeline_records_parse_failures_in_run_summary(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.progressive_queue_enabled = True
                config.cheap_ai_top_n = 1
                config.strong_ai_top_n = 1
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, APIs\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                sample_job = Job(
                    id="metrics-parse-1",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Build APIs with Python",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="jobspy",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(sample_job, {"sample": True})]
                pipeline.scorer.cheap_evaluate = lambda *args, **kwargs: StageEvaluation(
                    "cheap",
                    "openai",
                    "gpt-5-nano",
                    "error",
                    "error",
                    None,
                    None,
                    "Evaluation failed.",
                    [],
                    [],
                    10,
                    5,
                    0.0,
                    "invalid_json: Model response could not be parsed as JSON.",
                    datetime.now(timezone.utc).isoformat(),
                )
                result = pipeline.run()
                self.assertEqual(result.status, "completed")
                self.assertIn("cheap-stage scoring failures", result.message)
                latest_run = database.latest_run()
                self.assertIsNotNone(latest_run)
                assert latest_run is not None
                self.assertEqual(latest_run["cheap_parse_failures"], 1)
                self.assertEqual(latest_run["strong_shortlist_size"], 0)
                self.assertIn("cheap_stage_failed", latest_run["error_summary"])

    def test_semi_auto_approve_requires_generated_documents(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, SQL\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                sample_job = Job(
                    id="6",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Build APIs",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="jobspy",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(sample_job, {"sample": True})]
                pipeline.run()
                with self.assertRaisesRegex(ValueError, "Generate documents first"):
                    pipeline.approve_and_send("6")

    def test_generate_documents_for_job_creates_files(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, SQL\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                sample_job = Job(
                    id="7",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Build APIs",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="jobspy",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(sample_job, {"sample": True})]
                pipeline.run()
                docs = pipeline.generate_documents_for_job("7")
                self.assertTrue(docs.output_dir.exists())
                self.assertTrue(docs.resume_pdf_path.exists())
                self.assertTrue(docs.cover_letter_pdf_path.exists())
                self.assertTrue(docs.cover_letter_docx_path.exists())
                row = database.get_review_row("7")
                self.assertEqual(row["document_status"], "generated")
                self.assertTrue(row["generated_at"])

    def test_generate_documents_for_job_emits_progress_events(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, SQL\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                sample_job = Job(
                    id="70",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Build APIs",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="jobspy",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(sample_job, {"sample": True})]
                pipeline.run()

                def fake_generate(output_dir, job, resume, score, *, ai_notes="", tailoring_payload=None, progress_callback=None):
                    if progress_callback:
                        progress_callback("tailoring_resume", "Tailoring resume...", 3)
                        progress_callback("exporting_pdf", "Exporting PDF...", 4)
                        progress_callback("validating_pages", "Validating final page count...", 5)
                        progress_callback("writing_cover_letter", "Writing cover letter...", 6)
                    target_dir = output_dir / "Acme_2026-01-01_deadbeef"
                    target_dir.mkdir(parents=True, exist_ok=True)
                    resume_pdf = target_dir / "resume.pdf"
                    cover_txt = target_dir / "cover_letter.txt"
                    cover_docx = target_dir / "cover_letter.docx"
                    cover_pdf = target_dir / "cover_letter.pdf"
                    resume_pdf.write_text("pdf", encoding="utf-8")
                    cover_txt.write_text("cover", encoding="utf-8")
                    cover_docx.write_text("docx", encoding="utf-8")
                    cover_pdf.write_text("pdf", encoding="utf-8")
                    from doc_generator import GeneratedDocs
                    return GeneratedDocs(
                        output_dir=target_dir,
                        resume_pdf_path=resume_pdf,
                        cover_letter_pdf_path=cover_pdf,
                        cover_letter_docx_path=cover_docx,
                        cover_letter_txt_path=cover_txt,
                    )

                pipeline.doc_generator.generate = fake_generate
                events: list[tuple[str, str, int]] = []
                pipeline.generate_documents_for_job("70", progress_callback=lambda stage, message, progress: events.append((stage, message, progress)))
                self.assertEqual(
                    [stage for stage, _message, _progress in events],
                    ["starting", "loading_job", "parsing_resume", "tailoring_resume", "exporting_pdf", "validating_pages", "writing_cover_letter", "saving_database_state", "completed"],
                )

    def test_generate_documents_for_job_records_page_limit_failure(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, SQL\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                sample_job = Job(
                    id="71",
                    title="Director of Ecommerce",
                    employer="MILK BAR",
                    location="New York, NY",
                    salary_range="",
                    description_full="Lead ecommerce growth, reporting, and site optimization.",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="jobspy",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(sample_job, {"sample": True})]
                pipeline.run()

                from doc_generator import GeneratedDocs

                def fake_generate(output_dir, job, resume, score, *, ai_notes="", tailoring_payload=None, progress_callback=None):
                    if progress_callback:
                        progress_callback("tailoring_resume", "Tailoring resume...", 3)
                        progress_callback("exporting_pdf", "Exporting PDF...", 4)
                        progress_callback("validating_pages", "Validating final page count...", 5)
                    return GeneratedDocs(
                        output_dir=output_dir,
                        resume_pdf_path=Path(),
                        cover_letter_pdf_path=Path(),
                        resume_docx_path=Path(),
                        cover_letter_docx_path=Path(),
                        cover_letter_txt_path=Path(),
                        status="failed",
                        error_message="JobBot could not compress the tailored resume to 2 pages.",
                    )

                pipeline.doc_generator.generate = fake_generate
                with self.assertRaisesRegex(ValueError, "compress the tailored resume to 2 pages"):
                    pipeline.generate_documents_for_job("71")
                row = database.get_review_row("71")
                self.assertEqual(row["document_status"], "failed")
                self.assertIn("compress the tailored resume to 2 pages", row["document_error"].lower())

    def test_load_resume_reparses_legacy_cache_without_structured_sections(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "ROBERT THOM\nrobert@example.com | linkedin.com/in/example | (530) 220-4847\n"
                    "WORK EXPERIENCE\n"
                    "Underdog Strategies, New York, NY — Digital Advertising and Field Manager\n"
                    "JULY 2024 - Present\n"
                    "● Led digital ad strategy and reporting.\n"
                    "EDUCATION\n"
                    "University of California, Davis\n"
                    "Bachelor of Arts in Political Science Class of 2015\n"
                    "KEY SKILLS\n"
                    "Operations, Project Management, CRM, Data Analysis\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                paths.resume_json.write_text(
                    '{"source_path": "' + str(resume_path).replace("\\", "\\\\") + '", "raw_text": "legacy", "name": "ROBERT THOM", "email": "old@example.com", "phone": "555", "summary": "legacy", "skills": ["legacy"], "experience_lines": ["legacy"]}',
                    encoding="utf-8",
                )
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                resume = pipeline._load_resume()
                self.assertTrue(resume.header_lines)
                self.assertTrue(resume.education_lines)
                self.assertTrue(resume.work_experience_entries)

    def test_load_resume_reparses_docx_cache_without_template_metadata(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                from docx import Document

                paths = build_app_paths()
                config = load_or_create_config(paths)
                resume_path = Path(tmp) / "Robert Thom Resume 2026.docx"
                doc = Document()
                doc.add_paragraph("ROBERT THOM")
                doc.add_paragraph("robert@example.com | linkedin.com/in/example | (530) 220-4847")
                doc.add_paragraph("WORK EXPERIENCE", style="Heading 1")
                doc.add_paragraph("Underdog Strategies, New York, NY — Digital Advertising and Field Manager", style="Heading 2")
                doc.add_paragraph("JULY 2024 - Present")
                bullet = doc.add_paragraph(style="List Bullet")
                bullet.add_run("Led digital ad strategy and reporting")
                doc.add_paragraph("EDUCATION", style="Heading 1")
                doc.add_paragraph("University of California, Davis", style="Heading 2")
                doc.add_paragraph("KEY SKILLS", style="Heading 1")
                doc.add_paragraph("Operations, Project Management, CRM, Data Analysis")
                doc.save(str(resume_path))

                config.resume_source_path = str(resume_path)
                paths.resume_json.write_text(
                    (
                        '{'
                        f'"source_path": "{str(resume_path).replace("\\", "\\\\")}", '
                        '"raw_text": "legacy", '
                        '"name": "WORK EXPERIENCE", '
                        '"email": "old@example.com", '
                        '"phone": "555", '
                        '"summary": "legacy", '
                        '"skills": ["legacy"], '
                        '"experience_lines": ["legacy"], '
                        '"header_lines": ["robert@example.com | linkedin.com/in/example | (530) 220-4847"], '
                        '"education_lines": ["University of California, Davis"], '
                        '"key_skills_lines": ["Operations, Project Management, CRM, Data Analysis"], '
                        '"work_experience_entries": [{"role_line": "Underdog Strategies, New York, NY — Digital Advertising and Field Manager", "date_line": "JULY 2024 - Present", "bullets": ["Led digital ad strategy and reporting"]}]'
                        '}'
                    ),
                    encoding="utf-8",
                )
                database = Database(paths.database_file)
                database.initialize()
                pipeline = JobBotPipeline(config, paths, database)
                resume = pipeline._load_resume()
                self.assertEqual(resume.name, "ROBERT THOM")
                self.assertTrue(resume.key_skills_templates)
                self.assertTrue(resume.work_experience_entries[0].role_template is not None)
                self.assertTrue(resume.work_experience_entries[0].bullet_templates)

    def test_recent_duplicate_from_prior_day_window_is_skipped(self) -> None:
        with workspace_temp_dir() as tmp:
            with patch.dict(os.environ, {"APPDATA": str(tmp)}):
                paths = build_app_paths()
                config = load_or_create_config(paths)
                config.automation_mode = "semi_auto"
                config.skip_ai_scoring_in_semi_auto = True
                resume_path = Path(tmp) / "resume.txt"
                resume_path.write_text(
                    "Jane Candidate\njane@example.com\n555-123-4567\nSummary: Python engineer\nSkills: Python, SQL\nBuilt APIs\n",
                    encoding="utf-8",
                )
                config.resume_source_path = str(resume_path)
                database = Database(paths.database_file)
                database.initialize()
                existing_job = Job(
                    id="5",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Build APIs",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="jobspy",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at=datetime.now(timezone.utc).isoformat(),
                )
                database.upsert_job(existing_job, {"sample": True})
                pipeline = JobBotPipeline(config, paths, database)
                incoming_job = Job(
                    id="5",
                    title="Python Developer",
                    employer="Acme",
                    location="Remote",
                    salary_range="100-120",
                    description_full="Build APIs",
                    apply_method="board",
                    apply_url="https://example.com",
                    hiring_manager_email="",
                    source="jobspy",
                    posted_at="2026-01-01T00:00:00+00:00",
                    scraped_at="2026-01-01T00:00:00+00:00",
                )
                pipeline.scraper.fetch_jobs = lambda: [(incoming_job, {"sample": True})]
                result = pipeline.run()
                self.assertEqual(result.status, "completed")
                self.assertEqual(result.jobs_matched, 0)
